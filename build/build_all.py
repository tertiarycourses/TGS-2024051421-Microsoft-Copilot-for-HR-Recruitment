#!/usr/bin/env python3
"""Build aligned PPTX, Learner Guide, Lesson Plan and per-activity PDFs."""

from __future__ import annotations

import json
import math
import os
import re
from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches as DInches, Pt as DPt, RGBColor as DRGBColor
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_DATA_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn as p_qn
from pptx.util import Inches, Pt
from lxml import etree
from pypdf import PdfReader, PdfWriter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (BaseDocTemplate, Frame, PageTemplate, PageBreak,
                                Paragraph, Spacer, Table, TableStyle, KeepTogether)

from course_content import *

ROOT = Path(__file__).resolve().parents[1]
CW = ROOT / "courseware"
ACT = ROOT / "activities"
QA = ROOT / "qa"
ASSETS = CW / "assets"
CW.mkdir(exist_ok=True)
ACT.mkdir(exist_ok=True)
QA.mkdir(exist_ok=True)

BLUE = RGBColor(0x1F, 0x6F, 0xEB)
TEAL = RGBColor(0x10, 0x8A, 0x73)
VIOLET = RGBColor(0x6D, 0x3F, 0xD2)
AMBER = RGBColor(0xC7, 0x76, 0x00)
RED = RGBColor(0xC2, 0x41, 0x3A)
INK = RGBColor(0x16, 0x1B, 0x26)
GREY = RGBColor(0x5B, 0x63, 0x72)
LIGHT = RGBColor(0xF5, 0xF8, 0xFC)
LINE = RGBColor(0xD7, 0xE0, 0xEA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PALETTE = [BLUE, VIOLET, TEAL, AMBER]

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
section_slide_ids: set[int] = set()
slide_map = {"topics": {}, "activities": {}, "markers": {}}


def add_rect(slide, x, y, w, h, fill=LIGHT, line=None, radius=False):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, x, y, w, h)
    shape.fill.solid(); shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line; shape.line.width = Pt(1)
    return shape


def add_text(slide, x, y, w, h, text, size=14, color=INK, bold=False,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.04, font="Arial"):
    box = slide.shapes.add_textbox(x, y, w, h)
    box.text_frame.clear(); box.text_frame.word_wrap = True
    box.text_frame.margin_left = Inches(margin); box.text_frame.margin_right = Inches(margin)
    box.text_frame.margin_top = Inches(margin); box.text_frame.margin_bottom = Inches(margin)
    box.text_frame.vertical_anchor = valign
    p = box.text_frame.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = str(text); r.font.name = font; r.font.size = Pt(size)
    r.font.bold = bold; r.font.color.rgb = color
    return box


def fit_title_size(title):
    n = len(title)
    return 29 if n <= 52 else 26 if n <= 65 else 23 if n <= 82 else 21


def footer(slide, number):
    add_text(slide, Inches(0.72), Inches(6.96), Inches(7.2), Inches(0.25),
             f"{SHORT_TITLE} | {COURSE_CODE}", 8, GREY)
    add_text(slide, Inches(4.8), Inches(6.96), Inches(4.0), Inches(0.25),
             "(c) 2026 Tertiary Infotech Academy Pte Ltd", 8, GREY, align=PP_ALIGN.CENTER)
    add_text(slide, Inches(11.95), Inches(6.96), Inches(0.6), Inches(0.25),
             str(number), 8, GREY, align=PP_ALIGN.RIGHT)


def new_slide(title=None, kicker=None, accent=BLUE, numbered=True):
    slide = prs.slides.add_slide(BLANK)
    add_rect(slide, 0, 0, prs.slide_width, prs.slide_height, WHITE)
    if title is not None:
        add_rect(slide, 0, 0, Inches(0.22), Inches(1.48), accent)
        add_text(slide, Inches(0.72), Inches(0.34), Inches(11.85), Inches(0.3),
                 (kicker or "COURSE CONTENT").upper(), 12, accent, True)
        add_text(slide, Inches(0.72), Inches(0.69), Inches(11.85), Inches(0.62),
                 title, fit_title_size(title), INK, True)
        add_rect(slide, Inches(0.72), Inches(1.50), Inches(11.85), Inches(0.012), LINE)
    if numbered:
        footer(slide, len(prs.slides))
    return slide


def add_link_text(slide, x, y, w, h, label, url, size=13, color=BLUE):
    box = add_text(slide, x, y, w, h, label, size, color, True)
    run = box.text_frame.paragraphs[0].runs[0]
    run.hyperlink.address = url
    return box


def add_source(slide, source):
    label = re.sub(r"^https?://", "", source).split("/")[0]
    add_link_text(slide, Inches(8.2), Inches(6.68), Inches(4.3), Inches(0.2),
                  f"Source: {label}", source, 7.5, GREY)


def picture_fit(slide, path, x, y, w, h):
    path = str(path)
    with Image.open(path) as im:
        iw, ih = im.size
    box_ratio = w / h; im_ratio = iw / ih
    if im_ratio >= box_ratio:
        pw = w; ph = int(w / im_ratio); px = x; py = int(y + (h - ph) / 2)
    else:
        ph = h; pw = int(h * im_ratio); py = y; px = int(x + (w - pw) / 2)
    return slide.shapes.add_picture(path, px, py, width=pw, height=ph)


def picture_crop(slide, path, x, y, w, h, left=0.0, top=0.0, right=0.0, bottom=0.0):
    pic = slide.shapes.add_picture(str(path), x, y, width=w, height=h)
    pic.crop_left = left; pic.crop_top = top; pic.crop_right = right; pic.crop_bottom = bottom
    return pic


def cover_slide():
    slide = new_slide(numbered=False)
    add_rect(slide, 0, 0, prs.slide_width, Inches(0.18), BLUE)
    add_rect(slide, 0, Inches(7.32), prs.slide_width, Inches(0.18), TEAL)
    picture_fit(slide, ASSETS / "cover-interview-panel.png", Inches(5.55), Inches(0.18), Inches(7.78), Inches(7.14))
    add_rect(slide, Inches(0.0), Inches(0.18), Inches(6.25), Inches(7.14), WHITE)
    picture_fit(slide, ASSETS / "tertiary-infotech-logo.png", Inches(0.72), Inches(0.55), Inches(0.74), Inches(0.74))
    add_text(slide, Inches(0.72), Inches(1.60), Inches(5.0), Inches(0.35), "WSQ COURSE SLIDES", 13, BLUE, True)
    add_text(slide, Inches(0.72), Inches(2.05), Inches(5.1), Inches(1.65), TITLE, 34, INK, True)
    add_rect(slide, Inches(0.72), Inches(3.88), Inches(1.6), Inches(0.06), TEAL)
    add_text(slide, Inches(0.72), Inches(4.25), Inches(4.95), Inches(0.38), f"Course Code: {COURSE_CODE}", 15, INK, True)
    add_text(slide, Inches(0.72), Inches(4.70), Inches(4.95), Inches(0.70),
             f"{ORG}\nUEN {UEN}", 12.5, GREY)
    add_text(slide, Inches(0.72), Inches(5.65), Inches(4.95), Inches(0.34), f"Trainer: {TRAINER}", 12.5, GREY)
    add_text(slide, Inches(0.72), Inches(6.12), Inches(4.95), Inches(0.34),
             f"Version {VERSION} | {VERSION_DATE}", 12, GREY, True)
    add_text(slide, Inches(0.72), Inches(6.75), Inches(4.95), Inches(0.28),
             "Human judgement remains accountable for every hiring decision.", 10, TEAL, True)


def section_slide(topic_no, title, subtitle):
    slide = new_slide(numbered=True)
    section_slide_ids.add(len(prs.slides) - 1)
    add_rect(slide, Inches(0.72), Inches(1.90), Inches(11.85), Inches(3.60), LIGHT)
    add_rect(slide, Inches(0.72), Inches(1.90), Inches(0.14), Inches(3.60), PALETTE[(topic_no-1) % 4])
    add_text(slide, Inches(1.40), Inches(2.35), Inches(7.8), Inches(0.35), f"TOPIC {topic_no}", 15, PALETTE[(topic_no-1)%4], True)
    add_text(slide, Inches(1.40), Inches(2.85), Inches(9.5), Inches(1.10), title, 33, INK, True)
    add_text(slide, Inches(1.40), Inches(4.10), Inches(8.8), Inches(0.70), subtitle, 15, GREY)
    add_text(slide, Inches(10.20), Inches(2.20), Inches(1.7), Inches(1.2), f"{topic_no:02d}", 65, LINE, True, PP_ALIGN.RIGHT)


def cards_slide(title, cards, kicker, source=None, accent=BLUE, takeaway=None):
    slide = new_slide(title, kicker, accent)
    xs = [0.72, 3.72, 6.72, 9.72]
    for i, (head, body) in enumerate(cards[:4]):
        x = Inches(xs[i]); color = PALETTE[i]
        add_rect(slide, x, Inches(1.85), Inches(2.72), Inches(2.75), LIGHT)
        add_rect(slide, x, Inches(1.85), Inches(0.09), Inches(2.75), color)
        add_rect(slide, x+Inches(0.24), Inches(2.05), Inches(0.46), Inches(0.46), color, radius=True)
        add_text(slide, x+Inches(0.24), Inches(2.06), Inches(0.46), Inches(0.42), str(i+1), 15, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        add_text(slide, x+Inches(0.82), Inches(2.02), Inches(1.65), Inches(0.52), head, 13.5, INK, True)
        add_text(slide, x+Inches(0.28), Inches(2.72), Inches(2.18), Inches(1.45), body, 11.2, GREY)
        add_rect(slide, x+Inches(0.28), Inches(4.18), Inches(1.30), Inches(0.30), color, radius=True)
        add_text(slide, x+Inches(0.28), Inches(4.18), Inches(1.30), Inches(0.28), "EVIDENCE", 9, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    add_rect(slide, Inches(0.72), Inches(4.88), Inches(11.85), Inches(1.35), LIGHT)
    add_text(slide, Inches(1.00), Inches(5.04), Inches(2.1), Inches(0.28), "DECISION RULE", 11, accent, True)
    add_text(slide, Inches(1.00), Inches(5.43), Inches(11.20), Inches(0.54), takeaway or "Apply the rule consistently and retain the evidence used.", 13.5, INK, True)
    if source: add_source(slide, source)
    return slide


def control_slide(item, variant=0):
    if variant == 0:
        return cards_slide(item["title"] + " - controls and decision rules", item["controls"],
                           f'{item["codes"]} | CONTROL MATRIX', item["source"],
                           PALETTE[(item["topic"]-1)%4], item["rule"])
    slide = new_slide(item["title"] + " - controls and decision rules",
                      f'{item["codes"]} | CONTROL EVIDENCE', PALETTE[(item["topic"]-1)%4])
    if variant == 1:
        add_text(slide, Inches(0.92), Inches(1.78), Inches(4.0), Inches(0.30), "CONTROL TO APPLY", 11, BLUE, True)
        add_text(slide, Inches(7.08), Inches(1.78), Inches(4.0), Inches(0.30), "EVIDENCE TO RETAIN", 11, TEAL, True)
        for i, (head, body) in enumerate(item["controls"]):
            y = Inches(2.22 + i*0.86); color = PALETTE[i]
            add_rect(slide, Inches(0.92), y, Inches(5.15), Inches(0.66), LIGHT, LINE, True)
            add_text(slide, Inches(1.15), y+Inches(0.12), Inches(4.72), Inches(0.36), head, 12.2, INK, True)
            add_rect(slide, Inches(6.35), y+Inches(0.20), Inches(0.48), Inches(0.26), color, radius=True)
            add_text(slide, Inches(7.08), y+Inches(0.08), Inches(5.10), Inches(0.48), body, 11.5, GREY, True)
    elif variant == 2:
        add_rect(slide, Inches(0.82), Inches(1.84), Inches(5.65), Inches(3.82), RGBColor(0xFF,0xF4,0xF2), LINE, True)
        add_rect(slide, Inches(6.86), Inches(1.84), Inches(5.65), Inches(3.82), RGBColor(0xF1,0xFA,0xF7), LINE, True)
        add_text(slide, Inches(1.12), Inches(2.10), Inches(5.0), Inches(0.34), "RISK IF UNCONTROLLED", 12, RED, True)
        add_text(slide, Inches(7.16), Inches(2.10), Inches(5.0), Inches(0.34), "CONTROLLED PRACTICE", 12, TEAL, True)
        for i, (head, body) in enumerate(item["controls"]):
            add_text(slide, Inches(1.12), Inches(2.72+i*0.65), Inches(4.95), Inches(0.45), f"- {head}", 12, INK, True)
            add_text(slide, Inches(7.16), Inches(2.72+i*0.65), Inches(4.95), Inches(0.45), f"+ {body}", 11.3, INK, True)
    else:
        add_rect(slide, Inches(4.66), Inches(2.38), Inches(4.00), Inches(1.72), VIOLET, radius=True)
        add_text(slide, Inches(5.02), Inches(2.62), Inches(3.28), Inches(0.30), "DECISION GATE", 12, WHITE, True, PP_ALIGN.CENTER)
        add_text(slide, Inches(5.02), Inches(3.05), Inches(3.28), Inches(0.72), item["rule"], 11.2, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        positions = [(0.82,1.82),(8.96,1.82),(0.82,4.30),(8.96,4.30)]
        for i, ((head, body), (x,y)) in enumerate(zip(item["controls"], positions)):
            add_rect(slide, Inches(x), Inches(y), Inches(3.55), Inches(1.62), LIGHT, LINE, True)
            add_text(slide, Inches(x+0.22), Inches(y+0.18), Inches(3.10), Inches(0.32), head, 11.5, PALETTE[i], True)
            add_text(slide, Inches(x+0.22), Inches(y+0.60), Inches(3.10), Inches(0.70), body, 10.8, INK, True)
    add_rect(slide, Inches(0.82), Inches(5.93), Inches(11.68), Inches(0.38), TEAL, radius=True)
    add_text(slide, Inches(1.04), Inches(5.95), Inches(11.2), Inches(0.30), item["rule"], 10.8, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    add_source(slide, item["source"])
    return slide


def process_slide(concept_item, variant=0):
    slide = new_slide(concept_item["title"], f'{concept_item["codes"]} | MECHANISM', PALETTE[(concept_item["topic"]-1)%4])
    stages = concept_item["mechanism"]
    if variant == 1:
        add_text(slide, Inches(0.82), Inches(1.82), Inches(2.15), Inches(0.28), "INTERVIEWER LANE", 11, BLUE, True)
        add_text(slide, Inches(0.82), Inches(4.08), Inches(2.15), Inches(0.28), "CANDIDATE / EVIDENCE LANE", 11, TEAL, True)
        for i, stage in enumerate(stages):
            x = Inches(3.05 + (i % 2) * 4.70); y = Inches(1.80 + (i // 2) * 2.25)
            color = PALETTE[i]
            add_rect(slide, x, y, Inches(4.25), Inches(1.58), LIGHT, LINE, True)
            add_rect(slide, x+Inches(0.20), y+Inches(0.26), Inches(0.58), Inches(0.58), color, radius=True)
            add_text(slide, x+Inches(0.20), y+Inches(0.27), Inches(0.58), Inches(0.52), str(i+1), 18, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
            add_text(slide, x+Inches(0.98), y+Inches(0.24), Inches(2.95), Inches(0.78), stage, 15, INK, True, valign=MSO_ANCHOR.MIDDLE)
        add_rect(slide, Inches(0.82), Inches(5.92), Inches(11.65), Inches(0.42), TEAL, radius=True)
        add_text(slide, Inches(1.02), Inches(5.94), Inches(11.2), Inches(0.34), concept_item["rule"], 11.2, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        add_source(slide, concept_item["source"])
        return
    if variant == 2:
        positions = [(0.82, 2.05), (8.98, 2.05), (8.98, 4.30), (0.82, 4.30)]
        for i, (stage, (x, y)) in enumerate(zip(stages, positions)):
            color = PALETTE[i]
            add_rect(slide, Inches(x), Inches(y), Inches(3.50), Inches(1.48), LIGHT, LINE, True)
            add_text(slide, Inches(x+0.22), Inches(y+0.20), Inches(0.55), Inches(0.38), f"0{i+1}", 15, color, True)
            add_text(slide, Inches(x+0.88), Inches(y+0.20), Inches(2.40), Inches(0.86), stage, 13.5, INK, True, valign=MSO_ANCHOR.MIDDLE)
        add_rect(slide, Inches(4.86), Inches(2.82), Inches(3.58), Inches(2.18), VIOLET, radius=True)
        add_text(slide, Inches(5.06), Inches(3.05), Inches(3.18), Inches(0.30), "HUMAN JUDGEMENT", 11, WHITE, True, PP_ALIGN.CENTER)
        add_text(slide, Inches(5.06), Inches(3.60), Inches(3.18), Inches(0.90), concept_item["rule"], 11.5, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        add_source(slide, concept_item["source"])
        return
    if variant == 3:
        widths = [2.35, 2.95, 2.95, 2.35]
        x = 0.82
        for i, (stage, width) in enumerate(zip(stages, widths)):
            color = PALETTE[i]
            add_rect(slide, Inches(x), Inches(2.15), Inches(width), Inches(2.80), LIGHT, LINE, True)
            add_rect(slide, Inches(x), Inches(2.15), Inches(width), Inches(0.11), color)
            add_text(slide, Inches(x+0.22), Inches(2.52), Inches(width-0.44), Inches(0.35), ["INPUT", "TRANSFORM", "REVIEW", "RETAIN"][i], 11, color, True)
            add_text(slide, Inches(x+0.22), Inches(3.25), Inches(width-0.44), Inches(1.05), stage, 15, INK, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
            x += width + 0.18
        add_rect(slide, Inches(0.82), Inches(5.28), Inches(11.65), Inches(0.78), LIGHT, LINE, True)
        add_text(slide, Inches(1.08), Inches(5.47), Inches(11.1), Inches(0.36), f"Acceptance evidence: {concept_item['case'][3]}", 12.2, TEAL, True)
        add_source(slide, concept_item["source"])
        return
    xs = [0.85, 3.24, 5.64, 8.03]
    for i, stage in enumerate(stages):
        x = Inches(xs[i]); color = PALETTE[i]
        add_rect(slide, x, Inches(2.00), Inches(2.05), Inches(3.05), LIGHT)
        add_rect(slide, x, Inches(2.00), Inches(2.05), Inches(0.10), color)
        add_rect(slide, x+Inches(0.62), Inches(2.40), Inches(0.82), Inches(0.82), color, radius=True)
        add_text(slide, x+Inches(0.62), Inches(2.41), Inches(0.82), Inches(0.78), str(i+1), 28, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        add_text(slide, x+Inches(0.17), Inches(3.42), Inches(1.71), Inches(1.10), stage, 14, INK, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        if i < len(stages)-1:
            connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x+Inches(2.08), Inches(3.48), x+Inches(2.38), Inches(3.48))
            connector.line.color.rgb = color; connector.line.width = Pt(2)
            ln = connector._element.spPr.ln
            tail = OxmlElement('a:tailEnd'); tail.set('type', 'triangle'); ln.append(tail)
    add_rect(slide, Inches(10.43), Inches(2.00), Inches(2.05), Inches(3.05), LIGHT)
    add_rect(slide, Inches(10.43), Inches(2.00), Inches(2.05), Inches(0.10), RED)
    add_text(slide, Inches(10.70), Inches(2.40), Inches(1.50), Inches(0.32), "CONTROL POINT", 11, RED, True, PP_ALIGN.CENTER)
    add_text(slide, Inches(10.68), Inches(3.15), Inches(1.55), Inches(1.35), concept_item["rule"], 12, INK, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    add_rect(slide, Inches(0.85), Inches(5.35), Inches(11.63), Inches(0.82), LIGHT)
    add_text(slide, Inches(1.10), Inches(5.55), Inches(11.1), Inches(0.40), f"Observable output: {concept_item['case'][3]}", 12.5, TEAL, True)
    add_source(slide, concept_item["source"])


def chart_slide(item):
    cats, values, insight = item["chart"]
    slide = new_slide(item["title"] + " - editable evidence chart", f'{item["codes"]} | NATIVE POWERPOINT CHART', PALETTE[(item["topic"]-1)%4])
    data = ChartData(); data.categories = cats; data.add_series("Indicative value", values)
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.88), Inches(1.88), Inches(7.40), Inches(4.42), data).chart
    chart.has_title = False; chart.has_legend = False
    chart.value_axis.minimum_scale = 0
    vmax = max(values)
    if vmax <= 5:
        axis_max = 5
    elif vmax <= 10:
        axis_max = 10
    elif vmax <= 40:
        axis_max = math.ceil(vmax / 10) * 10
    elif vmax <= 60:
        axis_max = math.ceil(vmax / 10) * 10
    else:
        axis_max = 100
    chart.value_axis.maximum_scale = axis_max
    chart.value_axis.has_major_gridlines = True
    chart.category_axis.tick_labels.font.name = "Arial"; chart.category_axis.tick_labels.font.size = Pt(10)
    chart.value_axis.tick_labels.font.name = "Arial"; chart.value_axis.tick_labels.font.size = Pt(9)
    plot = chart.plots[0]; plot.has_data_labels = True
    plot.data_labels.position = XL_DATA_LABEL_POSITION.OUTSIDE_END
    plot.data_labels.font.name = "Arial"; plot.data_labels.font.size = Pt(10)
    for i, point in enumerate(plot.series[0].points):
        point.format.fill.solid(); point.format.fill.fore_color.rgb = PALETTE[i % len(PALETTE)]
    add_rect(slide, Inches(8.65), Inches(1.88), Inches(3.82), Inches(4.42), LIGHT)
    add_text(slide, Inches(8.95), Inches(2.15), Inches(3.20), Inches(0.30), "WHAT THE DATA SHOWS", 11.5, TEAL, True)
    add_text(slide, Inches(8.95), Inches(2.78), Inches(3.18), Inches(1.62), insight, 16, INK, True)
    add_text(slide, Inches(8.95), Inches(4.78), Inches(3.18), Inches(0.82), "Indicative teaching figures. Interpret the mechanism; do not treat these values as a universal benchmark.", 10.5, GREY)
    add_source(slide, item["source"])


def case_slide(item):
    special = None
    title_l = item["title"].lower()
    if "hr question generator" in title_l:
        special = ASSETS / "hr-question-generator.png"
    elif "ai interview practice lab" in title_l:
        special = ASSETS / "ai-interview-practice.png"
    elif "candidate ai feedback" in title_l:
        special = ASSETS / "candidate-ai-practice.png"
    elif "feedback should" in title_l:
        special = ASSETS / "evidence-feedback.png"
    slide = new_slide(item["title"] + " - worked evidence trace", f'{item["codes"]} | CASE AND EVIDENCE', PALETTE[(item["topic"]-1)%4])
    if special:
        add_rect(slide, Inches(0.72), Inches(1.82), Inches(6.65), Inches(4.68), LIGHT)
        picture_fit(slide, special, Inches(0.85), Inches(1.95), Inches(3.72), Inches(4.38))
        add_rect(slide, Inches(4.72), Inches(2.02), Inches(2.40), Inches(1.78), WHITE, LINE, True)
        picture_crop(slide, special, Inches(4.80), Inches(2.10), Inches(2.24), Inches(1.62), 0.0, 0.0, 0.48, 0.48)
        add_text(slide, Inches(4.80), Inches(3.88), Inches(2.24), Inches(0.25), "INPUT / RUBRIC ZOOM", 9.2, BLUE, True, PP_ALIGN.CENTER)
        add_rect(slide, Inches(4.72), Inches(4.28), Inches(2.40), Inches(1.78), WHITE, LINE, True)
        picture_crop(slide, special, Inches(4.80), Inches(4.36), Inches(2.24), Inches(1.62), 0.45, 0.38, 0.0, 0.0)
        x0 = Inches(7.70); w = Inches(4.87)
        labels = [("SCENARIO", item["case"][0]), ("AI OUTPUT / RISK", item["case"][1]),
                  ("HUMAN REVIEW", item["case"][2]), ("EVIDENCE", item["case"][3])]
        for i, (label, body) in enumerate(labels):
            y = Inches(1.83 + i*1.10); color = PALETTE[i]
            add_rect(slide, x0, y, w, Inches(0.94), LIGHT)
            add_rect(slide, x0, y, Inches(0.09), Inches(0.94), color)
            add_text(slide, x0+Inches(0.22), y+Inches(0.10), w-Inches(0.42), Inches(0.22), label, 9.5, color, True)
            add_text(slide, x0+Inches(0.22), y+Inches(0.37), w-Inches(0.42), Inches(0.48), body, 10.7, INK, True)
    else:
        labels = [("INPUT", item["case"][0]), ("AI OUTPUT / RISK", item["case"][1]),
                  ("HUMAN DECISION", item["case"][2]), ("AUDIT EVIDENCE", item["case"][3])]
        positions = [(0.72, 1.85), (6.72, 1.85), (0.72, 4.02), (6.72, 4.02)]
        for i, ((label, body), (x, y)) in enumerate(zip(labels, positions)):
            color = PALETTE[i]
            add_rect(slide, Inches(x), Inches(y), Inches(5.85), Inches(1.88), LIGHT)
            add_rect(slide, Inches(x), Inches(y), Inches(0.10), Inches(1.88), color)
            add_text(slide, Inches(x+0.30), Inches(y+0.18), Inches(5.15), Inches(0.28), label, 11, color, True)
            add_text(slide, Inches(x+0.30), Inches(y+0.70), Inches(5.15), Inches(0.82), body, 15, INK, True, valign=MSO_ANCHOR.MIDDLE)
    add_source(slide, item["source"])


def activity_slides(activity):
    start = len(prs.slides) + 1
    n = activity["num"]
    slide = new_slide(f'Activity {n}: {activity["title"]}', f'TOPIC {activity["topic"]} | ROLE PLAY AND PRACTICE', TEAL)
    add_rect(slide, Inches(0.72), Inches(1.84), Inches(1.58), Inches(0.46), TEAL, radius=True)
    add_text(slide, Inches(0.72), Inches(1.86), Inches(1.58), Inches(0.40), f"ACTIVITY {n}", 13, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    add_text(slide, Inches(10.55), Inches(1.84), Inches(2.02), Inches(0.46), f'{activity["duration"]} MIN', 13, TEAL, True, PP_ALIGN.RIGHT)
    add_text(slide, Inches(0.72), Inches(2.62), Inches(11.85), Inches(0.86), activity["scenario"], 18, INK, True)
    cards = [("ROLES", activity["roles"]), ("TOOLS", activity["tools"]), ("OUTPUT", activity["outcome"])]
    xs = [0.72, 4.72, 8.72]
    for i, (label, body) in enumerate(cards):
        add_rect(slide, Inches(xs[i]), Inches(4.00), Inches(3.72), Inches(1.78), LIGHT)
        add_rect(slide, Inches(xs[i]), Inches(4.00), Inches(3.72), Inches(0.10), PALETTE[i])
        add_text(slide, Inches(xs[i]+0.25), Inches(4.25), Inches(3.15), Inches(0.28), label, 11, PALETTE[i], True)
        add_text(slide, Inches(xs[i]+0.25), Inches(4.76), Inches(3.15), Inches(0.72), body, 12.5, INK, True)
    add_text(slide, Inches(0.72), Inches(6.34), Inches(11.85), Inches(0.28), f"Detailed procedure and acceptance checklist: Learner Guide and activities/activity-{n:02d}/", 10, GREY, True)

    workflows = {
        1:["Frame role", "Map competency", "Define evidence", "Assign panel", "Approve contract"],
        2:["Classify data", "Redact", "Threat-test", "Audit questions", "Record controls"],
        3:["Set inputs", "Generate pack", "Inspect gaps", "Review questions", "Approve guide"],
        4:["Select incident", "Draft item", "Create probes", "Anchor 1/3/5", "Pilot"],
        5:["Open", "Ask core", "Probe", "Close", "Debrief"],
        6:["Choose role", "Review rubric", "Interview", "Read evidence", "Retry"],
        7:["Check access", "Set cues", "Simulate failure", "Recover", "Document"],
        8:["Highlight evidence", "Score alone", "Reveal", "Calibrate", "Record"],
        9:["Select evidence", "State impact", "Suggest action", "Confirm", "Observe"],
        10:["Prepare", "Conduct", "Score", "Feedback", "Reflect"],
    }
    slide = new_slide(f'Activity {n} - mechanism and role hand-offs', f'TOPIC {activity["topic"]} | HIGH-LEVEL WORKFLOW', TEAL)
    stages = workflows[n]
    if n % 2:
        for i, stage in enumerate(stages):
            x = Inches(0.85 + i*2.39); color = PALETTE[i%4]
            add_rect(slide, x, Inches(2.30), Inches(2.05), Inches(2.70), LIGHT)
            add_rect(slide, x, Inches(2.30), Inches(2.05), Inches(0.10), color)
            add_rect(slide, x+Inches(0.62), Inches(2.68), Inches(0.82), Inches(0.82), color, radius=True)
            add_text(slide, x+Inches(0.62), Inches(2.69), Inches(0.82), Inches(0.78), str(i+1), 28, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
            add_text(slide, x+Inches(0.16), Inches(3.77), Inches(1.73), Inches(0.68), stage, 14, INK, True, PP_ALIGN.CENTER)
            if i < 4:
                conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x+Inches(2.05), Inches(3.60), x+Inches(2.34), Inches(3.60))
                conn.line.color.rgb = color; conn.line.width = Pt(2)
                tail = OxmlElement('a:tailEnd'); tail.set('type', 'triangle'); conn._element.spPr.ln.append(tail)
    else:
        add_text(slide, Inches(0.90), Inches(1.86), Inches(2.05), Inches(0.28), "PERFORM", 11, BLUE, True)
        add_text(slide, Inches(0.90), Inches(4.18), Inches(2.05), Inches(0.28), "REVIEW + RETAIN", 11, TEAL, True)
        positions=[(3.05,1.84),(7.70,1.84),(3.05,4.12),(7.70,4.12),(10.12,2.98)]
        for i,(stage,(x,y)) in enumerate(zip(stages,positions)):
            color=PALETTE[i%4]
            add_rect(slide, Inches(x), Inches(y), Inches(2.20), Inches(1.20), LIGHT, LINE, True)
            add_text(slide, Inches(x+0.18), Inches(y+0.14), Inches(0.40), Inches(0.28), str(i+1), 13, color, True)
            add_text(slide, Inches(x+0.63), Inches(y+0.15), Inches(1.38), Inches(0.68), stage, 12.5, INK, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    add_rect(slide, Inches(0.85), Inches(5.40), Inches(11.63), Inches(0.72), LIGHT)
    add_text(slide, Inches(1.10), Inches(5.58), Inches(11.0), Inches(0.32), "This slide shows control hand-offs only. Full click paths, scripts and recovery steps stay in the Learner Guide.", 11.8, TEAL, True)

    slide = new_slide(f'Activity {n} - evidence and acceptance', f'TOPIC {activity["topic"]} | VERIFY', TEAL)
    checks = activity["checklist"][:5]
    if n % 2:
        for i, check in enumerate(checks):
            y = Inches(1.85 + i*0.90); color = PALETTE[i%4]
            add_rect(slide, Inches(0.72), y, Inches(11.85), Inches(0.70), LIGHT)
            add_rect(slide, Inches(0.92), y+Inches(0.13), Inches(0.44), Inches(0.44), TEAL, radius=True)
            add_text(slide, Inches(0.92), y+Inches(0.13), Inches(0.44), Inches(0.40), "OK", 9, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
            add_text(slide, Inches(1.60), y+Inches(0.15), Inches(10.55), Inches(0.38), check, 14, INK, True)
    else:
        for i,check in enumerate(checks):
            col=i%2; row=i//2; x=0.82+col*5.88; y=1.85+row*1.28
            add_rect(slide, Inches(x), Inches(y), Inches(5.55), Inches(1.02), LIGHT, LINE, True)
            add_rect(slide, Inches(x+0.20), Inches(y+0.26), Inches(0.48), Inches(0.48), TEAL, radius=True)
            add_text(slide, Inches(x+0.20), Inches(y+0.27), Inches(0.48), Inches(0.42), "OK", 9, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
            add_text(slide, Inches(x+0.88), Inches(y+0.18), Inches(4.35), Inches(0.62), check, 12.2, INK, True, valign=MSO_ANCHOR.MIDDLE)
    add_rect(slide, Inches(0.72), Inches(6.42), Inches(11.85), Inches(0.28), TEAL)
    add_text(slide, Inches(0.90), Inches(6.42), Inches(11.4), Inches(0.25), f"Evidence pack: {activity['outcome']}", 9.5, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    slide_map["activities"][str(n)] = {"title": activity["title"], "start": start, "end": len(prs.slides)}


def admin_cards(title, cards, kicker="COURSE ADMIN", accent=BLUE):
    cards_slide(title, cards, kicker, accent=accent, takeaway="Use this page as an operational reference during the course.")


def attendance_slide(end=False):
    title = "Digital Attendance (Mandatory)" if end else "Digital Attendance"
    cards = [("AM", "Scan the SSG/TRAQOM QR code and submit attendance."),
             ("PM", "Repeat after lunch; verify successful submission."),
             ("ASSESSMENT", "Complete the assessment digital attendance before papers."),
             ("SUPPORT", "Tell the trainer immediately if the submission fails.")]
    cards_slide(title, cards, "TRAQOM | SSG", source=LMS_URL, accent=TEAL,
                takeaway="Attendance evidence is mandatory; a screenshot of a QR code is not proof of submission.")


def flow_slide():
    slide = new_slide("Assessment Flow", "WSQ ASSESSMENT | ROLE PLAY", VIOLET)
    steps = ["TRAQOM", "Assessment attendance", "WA then Role Play", "Submit on LMS", "Sign summary record"]
    for i, step in enumerate(steps):
        x = Inches(0.85 + i*2.39); color = PALETTE[i%4]
        add_rect(slide, x, Inches(2.25), Inches(2.05), Inches(3.05), LIGHT)
        add_rect(slide, x, Inches(2.25), Inches(2.05), Inches(0.10), color)
        add_rect(slide, x+Inches(0.62), Inches(2.67), Inches(0.82), Inches(0.82), color, radius=True)
        add_text(slide, x+Inches(0.62), Inches(2.68), Inches(0.82), Inches(0.78), str(i+1), 28, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        add_text(slide, x+Inches(0.16), Inches(3.85), Inches(1.73), Inches(0.88), step, 13.5, INK, True, PP_ALIGN.CENTER)
        if i < 4:
            conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x+Inches(2.05), Inches(3.60), x+Inches(2.34), Inches(3.60))
            conn.line.color.rgb = color; conn.line.width = Pt(2)
            tail = OxmlElement('a:tailEnd'); tail.set('type', 'triangle'); conn._element.spPr.ln.append(tail)
    add_link_text(slide, Inches(3.30), Inches(5.85), Inches(6.7), Inches(0.34), LMS_URL, LMS_URL, 12, BLUE)


def trainer_profile_slide(name, credentials, bio_lines, accent=VIOLET, template=False):
    """House-style trainer PROFILE CARD (avatar, name, credentials, bio) - not a bullet grid."""
    kicker = "TRAINER PROFILE | TEMPLATE" if template else "TRAINER PROFILE"
    title = "About the Trainer" if template else f"About the Trainer - {name}"
    slide = new_slide(title, kicker, accent)
    add_rect(slide, Inches(0.72), Inches(1.78), Inches(11.85), Inches(4.55), LIGHT)
    add_rect(slide, Inches(0.72), Inches(1.78), Inches(0.13), Inches(4.55), accent)
    # avatar
    add_rect(slide, Inches(1.32), Inches(2.32), Inches(2.30), Inches(2.30), WHITE, LINE, True)
    add_rect(slide, Inches(1.62), Inches(2.62), Inches(1.70), Inches(1.70), accent, radius=True)
    initials = "?" if template else "".join(w[0] for w in name.replace("Dr ", "").split()[:2]).upper()
    add_text(slide, Inches(1.62), Inches(2.66), Inches(1.70), Inches(1.62), initials, 40, WHITE, True,
             PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    add_text(slide, Inches(1.32), Inches(4.74), Inches(2.30), Inches(0.26),
             "PHOTOGRAPH" if template else "TRAINER", 9, GREY, True, PP_ALIGN.CENTER)
    # name + credentials
    add_text(slide, Inches(4.10), Inches(2.32), Inches(8.10), Inches(0.52), name, 24, INK, True)
    add_rect(slide, Inches(4.10), Inches(2.92), Inches(1.55), Inches(0.05), accent)
    add_text(slide, Inches(4.10), Inches(3.06), Inches(8.10), Inches(0.34), credentials, 12.5, accent, True)
    for i, line in enumerate(bio_lines):
        add_text(slide, Inches(4.10), Inches(3.56 + i*0.46), Inches(8.10), Inches(0.42),
                 chr(8226) + "  " + line, 12, INK)
    add_text(slide, Inches(0.72), Inches(6.44), Inches(11.85), Inches(0.28),
             "Complete this profile with the assigned trainer's details before delivery." if template
             else "Ask the trainer for context, worked examples and industry practice at any point.",
             10.5, GREY, True, PP_ALIGN.CENTER)


def practice_exam_slide():
    """Practice Exam slide (house standard A4) - exams.tertiaryinfotech.com."""
    slide = new_slide("Practice Exam", "PREPARE | EXAMS.TERTIARYINFOTECH.COM", AMBER)
    add_rect(slide, Inches(0.72), Inches(1.82), Inches(7.05), Inches(4.05), LIGHT)
    add_rect(slide, Inches(0.72), Inches(1.82), Inches(7.05), Inches(0.10), AMBER)
    # browser mock
    add_rect(slide, Inches(1.02), Inches(2.16), Inches(6.45), Inches(3.35), WHITE, LINE, True)
    add_rect(slide, Inches(1.02), Inches(2.16), Inches(6.45), Inches(0.42), RGBColor(0xEC, 0xF1, 0xF7))
    for i, c in enumerate([RGBColor(0xE0,0x6C,0x62), RGBColor(0xE8,0xB4,0x39), RGBColor(0x63,0xB3,0x6B)]):
        add_rect(slide, Inches(1.18 + i*0.22), Inches(2.30), Inches(0.14), Inches(0.14), c, radius=True)
    add_rect(slide, Inches(1.92), Inches(2.26), Inches(5.35), Inches(0.24), WHITE, LINE, True)
    add_text(slide, Inches(2.02), Inches(2.26), Inches(5.15), Inches(0.24), PRACTICE_EXAM_URL, 9, BLUE, False, valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, Inches(1.28), Inches(2.78), Inches(5.95), Inches(0.36), "Generative AI for Interviewing - Practice Exam", 14, INK, True)
    add_text(slide, Inches(1.28), Inches(3.20), Inches(5.95), Inches(0.30), "Multiple attempts | Instant feedback | Same K1-K6 coverage as the WA", 10.5, GREY)
    for i, q in enumerate(["Q1  Which control must be applied before prompting with a CV?",
                           "Q2  A rater has no evidence for a competency. What is recorded?",
                           "Q3  Identify the defect in this interview question."]):
        y = 3.66 + i*0.52
        add_rect(slide, Inches(1.28), Inches(y), Inches(5.95), Inches(0.44), WHITE, LINE, True)
        add_text(slide, Inches(1.44), Inches(y), Inches(5.65), Inches(0.44), q, 10, INK, valign=MSO_ANCHOR.MIDDLE)
    # right panel
    add_rect(slide, Inches(8.05), Inches(1.82), Inches(4.52), Inches(4.05), WHITE, LINE, True)
    add_text(slide, Inches(8.32), Inches(2.06), Inches(4.0), Inches(0.30), "HOW TO USE IT", 11, AMBER, True)
    for i, step in enumerate(["Open the practice exam before the written assessment.",
                              "Attempt it closed-book to find your real gaps.",
                              "Review the slides and Learner Guide for each miss.",
                              "Re-attempt until you are consistently correct."]):
        y = 2.48 + i*0.72
        add_rect(slide, Inches(8.32), Inches(y), Inches(0.42), Inches(0.42), PALETTE[i % 4], radius=True)
        add_text(slide, Inches(8.32), Inches(y), Inches(0.42), Inches(0.42), str(i+1), 13, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        add_text(slide, Inches(8.90), Inches(y-0.02), Inches(3.45), Inches(0.50), step, 11, INK, valign=MSO_ANCHOR.MIDDLE)
    add_rect(slide, Inches(8.32), Inches(5.42), Inches(3.95), Inches(0.34), AMBER, radius=True)
    add_link_text(slide, Inches(8.42), Inches(5.44), Inches(3.75), Inches(0.30), "Open the practice exam", PRACTICE_EXAM_URL, 11, WHITE)
    add_text(slide, Inches(0.72), Inches(6.10), Inches(11.85), Inches(0.30),
             "The practice exam is formative: it prepares you for the WA, it is not the WA.", 11.5, GREY, True, PP_ALIGN.CENTER)


def lab_login_slide():
    """Opening sign-in slide. The password is NEVER printed - the trainer says it in class."""
    slide = new_slide("Sign In to Your Lab Account", "MICROSOFT 365 COPILOT | LAB ACCESS", BLUE)
    add_rect(slide, Inches(0.72), Inches(1.78), Inches(5.75), Inches(3.05), LIGHT)
    add_rect(slide, Inches(0.72), Inches(1.78), Inches(5.75), Inches(0.10), BLUE)
    add_text(slide, Inches(0.95), Inches(2.00), Inches(5.3), Inches(0.30), "STEP 1  OPEN THE PORTAL", 11, BLUE, True)
    add_link_text(slide, Inches(0.95), Inches(2.34), Inches(5.3), Inches(0.32), M365_PORTAL, M365_PORTAL, 13, BLUE)
    add_text(slide, Inches(0.95), Inches(2.80), Inches(5.3), Inches(0.30), "STEP 2  SIGN IN WITH YOUR TRAINING ACCOUNT", 11, BLUE, True)
    add_text(slide, Inches(0.95), Inches(3.14), Inches(5.3), Inches(0.30), LAB_LOGIN_USER_1, 12.5, INK, True)
    add_text(slide, Inches(0.95), Inches(3.46), Inches(5.3), Inches(0.30), LAB_LOGIN_USER_2, 12.5, INK, True)
    add_rect(slide, Inches(0.95), Inches(3.88), Inches(5.28), Inches(0.62), WHITE, AMBER, True)
    add_text(slide, Inches(1.12), Inches(3.98), Inches(4.95), Inches(0.45), LAB_LOGIN_NOTE, 11.5, INK, True)

    add_rect(slide, Inches(6.82), Inches(1.78), Inches(5.75), Inches(3.05), LIGHT)
    add_rect(slide, Inches(6.82), Inches(1.78), Inches(5.75), Inches(0.10), TEAL)
    add_text(slide, Inches(7.05), Inches(2.00), Inches(5.3), Inches(0.30), "WHAT YOU WILL USE TODAY", 11, TEAL, True)
    for i, line in enumerate([
            "Copilot Chat - prompt engineering for interviews",
            "Copilot Studio - build and publish your own agent",
            "SharePoint - the synthetic candidate and policy corpus",
            "Five reference agents, already published for you"]):
        add_text(slide, Inches(7.05), Inches(2.38 + i*0.52), Inches(5.3), Inches(0.46), chr(8226) + "  " + line, 12, INK)

    add_rect(slide, Inches(0.72), Inches(5.02), Inches(11.85), Inches(1.12), WHITE, RED, True)
    add_text(slide, Inches(0.95), Inches(5.14), Inches(11.4), Inches(0.30), "BEFORE YOU PROMPT ANYTHING", 11, RED, True)
    add_text(slide, Inches(0.95), Inches(5.46), Inches(11.4), Inches(0.58),
             "Use only the supplied synthetic data. Never paste NRIC, date of birth, photograph, address, nationality, "
             "marital or family status, health information, or any real candidate's CV into an AI tool.", 12, INK)
    add_text(slide, Inches(0.72), Inches(6.34), Inches(11.85), Inches(0.30),
             "Sign in first; every activity today depends on this account.", 11.5, GREY, True, PP_ALIGN.CENTER)


def copilot_env_slide():
    """The dedicated course environment, agents and workflows."""
    slide = new_slide("Your Copilot Studio Environment", "COPILOT STUDIO | COURSE ENVIRONMENT", VIOLET)
    add_rect(slide, Inches(0.72), Inches(1.78), Inches(11.85), Inches(0.86), LIGHT)
    add_rect(slide, Inches(0.72), Inches(1.78), Inches(11.85), Inches(0.10), VIOLET)
    add_text(slide, Inches(0.95), Inches(1.98), Inches(11.4), Inches(0.30), "ENVIRONMENT", 10.5, VIOLET, True)
    add_text(slide, Inches(0.95), Inches(2.24), Inches(11.4), Inches(0.32), COPILOT_STUDIO_ENV, 14, INK, True)
    rows = [(n, d) for n, d, _u in COPILOT_AGENTS]
    for i, (name, desc) in enumerate(rows):
        y = 2.86 + i*0.66
        add_rect(slide, Inches(0.72), Inches(y), Inches(11.85), Inches(0.58), WHITE, LINE, True)
        add_rect(slide, Inches(0.72), Inches(y), Inches(0.09), Inches(0.58), PALETTE[i % 4])
        add_text(slide, Inches(0.98), Inches(y+0.05), Inches(3.55), Inches(0.48), name, 11.5, INK, True, valign=MSO_ANCHOR.MIDDLE)
        add_text(slide, Inches(4.60), Inches(y+0.05), Inches(7.85), Inches(0.48), desc, 10.5, GREY, valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, Inches(0.72), Inches(6.30), Inches(11.85), Inches(0.30),
             "All five agents are published. Workflows in this environment are named DO NOT DELETE - they are shared class assets.",
             11, GREY, True, PP_ALIGN.CENTER)


def sharepoint_corpus_slide():
    slide = new_slide("The Practice Data: SharePoint Hiring Sites", "SHAREPOINT | SYNTHETIC CORPUS", TEAL)
    add_rect(slide, Inches(0.72), Inches(1.78), Inches(11.85), Inches(0.80), LIGHT)
    add_rect(slide, Inches(0.72), Inches(1.78), Inches(11.85), Inches(0.10), TEAL)
    add_text(slide, Inches(0.95), Inches(2.02), Inches(11.4), Inches(0.42), SHAREPOINT_CORPUS + "  All records are invented for training.", 12.5, INK, True)
    for i, (name, sector, _url) in enumerate(SHAREPOINT_SITES):
        y = 2.80 + i*0.64
        add_rect(slide, Inches(0.72), Inches(y), Inches(11.85), Inches(0.56), WHITE, LINE, True)
        add_rect(slide, Inches(0.72), Inches(y), Inches(0.09), Inches(0.56), PALETTE[i % 4])
        add_text(slide, Inches(0.98), Inches(y+0.04), Inches(5.20), Inches(0.48), name, 11.5, INK, True, valign=MSO_ANCHOR.MIDDLE)
        add_text(slide, Inches(6.30), Inches(y+0.04), Inches(2.60), Inches(0.48), sector, 10.5, GREY, valign=MSO_ANCHOR.MIDDLE)
        add_text(slide, Inches(9.10), Inches(y+0.04), Inches(3.35), Inches(0.48), "Resumes | HR policies | JD", 10.5, GREY, valign=MSO_ANCHOR.MIDDLE)
    add_rect(slide, Inches(0.72), Inches(6.06), Inches(11.85), Inches(0.62), WHITE, RED, True)
    add_text(slide, Inches(0.95), Inches(6.16), Inches(11.4), Inches(0.44),
             "Every resume is synthetic. One of them hides a prompt-injection instruction - your screening agent must catch it, not obey it.",
             11.5, INK, True)


def prompt_pattern_slide():
    slide = new_slide("The Four-Part Prompt Pattern", "PROMPT ENGINEERING | CORE METHOD", BLUE)
    for i, (part, desc) in enumerate(PROMPT_PATTERN):
        x = Inches(0.72 + i*2.99)
        add_rect(slide, x, Inches(2.05), Inches(2.78), Inches(2.62), LIGHT)
        add_rect(slide, x, Inches(2.05), Inches(2.78), Inches(0.10), PALETTE[i % 4])
        add_rect(slide, x+Inches(1.09), Inches(2.36), Inches(0.62), Inches(0.62), PALETTE[i % 4], radius=True)
        add_text(slide, x+Inches(1.09), Inches(2.37), Inches(0.62), Inches(0.60), str(i+1), 20, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        add_text(slide, x+Inches(0.16), Inches(3.12), Inches(2.46), Inches(0.34), part, 13.5, INK, True, PP_ALIGN.CENTER)
        add_text(slide, x+Inches(0.16), Inches(3.50), Inches(2.46), Inches(1.02), desc, 10.5, GREY, align=PP_ALIGN.CENTER)
    add_rect(slide, Inches(0.72), Inches(4.94), Inches(11.85), Inches(1.16), WHITE, LINE, True)
    add_text(slide, Inches(0.95), Inches(5.04), Inches(11.4), Inches(0.28), "WEAK PROMPT", 10.5, RED, True)
    add_text(slide, Inches(0.95), Inches(5.30), Inches(11.4), Inches(0.30), "\"Give me interview questions for a data analyst.\"", 12, GREY)
    add_text(slide, Inches(0.95), Inches(5.62), Inches(11.4), Inches(0.28), "STRONG PROMPT", 10.5, TEAL, True)
    add_text(slide, Inches(0.95), Inches(5.86), Inches(11.4), Inches(0.30),
             "Role + the approved JD + one task + constraints (no protected traits, one competency per question, anchors required).", 11.5, INK)
    add_text(slide, Inches(0.72), Inches(6.26), Inches(11.85), Inches(0.30),
             "Most bad AI output is a good model answering a bad prompt.", 11.5, GREY, True, PP_ALIGN.CENTER)


def tool_screenshot_slide(title, image, url, points):
    slide = new_slide(title, "COURSE AI TOOL | AUTHENTIC UI", TEAL)
    add_rect(slide, Inches(0.72), Inches(1.82), Inches(7.25), Inches(4.75), LIGHT)
    picture_fit(slide, image, Inches(0.85), Inches(1.95), Inches(4.10), Inches(4.48))
    add_rect(slide, Inches(5.10), Inches(1.98), Inches(2.55), Inches(2.05), WHITE, LINE, True)
    picture_crop(slide, image, Inches(5.18), Inches(2.06), Inches(2.39), Inches(1.89), 0.0, 0.0, 0.48, 0.48)
    add_text(slide, Inches(5.20), Inches(4.15), Inches(2.35), Inches(0.28), "ZOOM: INPUTS", 9.5, BLUE, True, PP_ALIGN.CENTER)
    add_rect(slide, Inches(5.10), Inches(4.52), Inches(2.55), Inches(1.70), WHITE, LINE, True)
    picture_crop(slide, image, Inches(5.18), Inches(4.60), Inches(2.39), Inches(1.54), 0.45, 0.38, 0.0, 0.0)
    for i, point in enumerate(points):
        y = Inches(1.90 + i*1.05); color = PALETTE[i%4]
        add_rect(slide, Inches(8.32), y, Inches(4.25), Inches(0.86), LIGHT)
        add_rect(slide, Inches(8.52), y+Inches(0.20), Inches(0.44), Inches(0.44), color, radius=True)
        add_text(slide, Inches(8.52), y+Inches(0.20), Inches(0.44), Inches(0.40), str(i+1), 13, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        add_text(slide, Inches(9.18), y+Inches(0.16), Inches(3.05), Inches(0.52), point, 11.5, INK, True)
    add_link_text(slide, Inches(8.32), Inches(6.15), Inches(4.25), Inches(0.30), url, url, 10.5, BLUE)


def lms_orientation_slide():
    slide = new_slide("Download Course Material on the LMS", "COURSE ACCESS | EDITABLE ORIENTATION MOCK", TEAL)
    add_rect(slide, Inches(0.72), Inches(1.82), Inches(8.10), Inches(4.78), WHITE, LINE, True)
    add_rect(slide, Inches(0.72), Inches(1.82), Inches(8.10), Inches(0.48), RGBColor(0xE9,0xEF,0xF7))
    for i,c in enumerate([RED, AMBER, TEAL]): add_rect(slide, Inches(0.94+i*0.25), Inches(1.97), Inches(0.12), Inches(0.12), c, radius=True)
    add_rect(slide, Inches(1.83), Inches(1.94), Inches(5.85), Inches(0.22), WHITE, LINE, True)
    add_text(slide, Inches(2.02), Inches(1.94), Inches(5.45), Inches(0.18), "lms-tms.tertiaryinfotech.com / course / TGS-2024051421", 8.2, GREY)
    add_rect(slide, Inches(0.96), Inches(2.55), Inches(1.55), Inches(3.67), INK, radius=True)
    add_text(slide, Inches(1.18), Inches(2.82), Inches(1.10), Inches(0.32), "MY COURSE", 10, WHITE, True)
    for i,label in enumerate(["Overview","Materials","Activities","Assessment"]):
        color = TEAL if label == "Materials" else WHITE
        add_text(slide, Inches(1.18), Inches(3.48+i*0.54), Inches(1.08), Inches(0.30), label, 9.5, color, label=="Materials")
    add_text(slide, Inches(2.88), Inches(2.65), Inches(5.45), Inches(0.42), "Generative AI for Interviewing", 19, INK, True)
    add_text(slide, Inches(2.88), Inches(3.10), Inches(5.45), Inches(0.28), COURSE_CODE, 10.5, GREY, True)
    materials = [("Slides", "PPTX + PDF"), ("Learner Guide", "DOCX + PDF"), ("Lesson Plan", "Trainer reference"), ("Activities", "10 folders")]
    for i,(name,meta) in enumerate(materials):
        y = Inches(3.64+i*0.57)
        add_rect(slide, Inches(2.88), y, Inches(5.35), Inches(0.44), LIGHT, LINE, True)
        add_text(slide, Inches(3.10), y+Inches(0.08), Inches(2.35), Inches(0.24), name, 10.5, INK, True)
        add_text(slide, Inches(5.40), y+Inches(0.08), Inches(1.45), Inches(0.24), meta, 9.3, GREY)
        add_rect(slide, Inches(7.15), y+Inches(0.08), Inches(0.82), Inches(0.27), BLUE, radius=True)
        add_text(slide, Inches(7.15), y+Inches(0.08), Inches(0.82), Inches(0.24), "OPEN", 8.5, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    for i,(head,body) in enumerate([("1  SIGN IN","Use your assigned LMS account."),("2  OPEN COURSE","Match the title and course code."),("3  USE MATERIALS","Download learning files; upload assessment only in its area.")]):
        y=Inches(1.90+i*1.42)
        add_rect(slide, Inches(9.18), y, Inches(3.38), Inches(1.16), LIGHT, LINE, True)
        add_text(slide, Inches(9.43), y+Inches(0.18), Inches(2.90), Inches(0.28), head, 10.5, PALETTE[i], True)
        add_text(slide, Inches(9.43), y+Inches(0.53), Inches(2.90), Inches(0.45), body, 10.4, INK, True)
    add_link_text(slide, Inches(9.18), Inches(6.20), Inches(3.38), Inches(0.28), "Open LMS-TMS", LMS_URL, 10.5, BLUE)


def source_slides():
    items = list(SOURCES.items())
    chunk = 7
    for page, start in enumerate(range(0, len(items), chunk), 1):
        slide = new_slide(f"Sources and further reading ({page})", "REFERENCE | ACCESSED AUGUST 2026", BLUE)
        subset = items[start:start+chunk]
        for i, (key, url) in enumerate(subset):
            y = Inches(1.82 + i*0.67); color = PALETTE[i%4]
            add_rect(slide, Inches(0.72), y, Inches(11.85), Inches(0.52), LIGHT)
            add_rect(slide, Inches(0.72), y, Inches(0.08), Inches(0.52), color)
            add_text(slide, Inches(0.98), y+Inches(0.08), Inches(2.30), Inches(0.26), key.replace("_", " ").upper(), 9.5, color, True)
            domain = re.sub(r"^www\.", "", re.sub(r"^https?://", "", url).split("/")[0])
            add_link_text(slide, Inches(3.10), y+Inches(0.08), Inches(9.10), Inches(0.30), f"{domain} — open reference", url, 10.5, BLUE)


def _transition(slide, kind="fade", speed="fast"):
    sld = slide._element
    for old in sld.findall(p_qn("p:transition")): sld.remove(old)
    tr = etree.SubElement(sld, p_qn("p:transition")); tr.set("spd", speed); tr.set("advClick", "1")
    if kind == "push": etree.SubElement(tr, p_qn("p:push")).set("dir", "l")
    else: etree.SubElement(tr, p_qn("p:fade"))
    sld.append(tr)


def build_deck():
    cover_slide()
    attendance_slide()
    trainer_profile_slide("Trainer Name", "Qualification | Certification | Years of experience",
                          ["Industry background relevant to this course.",
                           "Professional certifications and teaching experience.",
                           "Areas of specialisation.",
                           "Contact: enquiry@tertiaryinfotech.com"], accent=BLUE, template=True)
    trainer_profile_slide(TRAINER, "PhD | Founder, Tertiary Infotech Academy | 15+ years in AI and automation",
                          ["Founder and technical trainer at Tertiary Infotech Academy.",
                           "Specialises in AI, cloud, automation and applied workplace learning.",
                           "Teaches mechanism-led, evidence-based practice, not tool tours.",
                           "Contact: enquiry@tertiaryinfotech.com"], accent=VIOLET)
    admin_cards("Learner Introduction", [("YOUR ROLE", "Interviewer, interviewee or both?"), ("ONE CHALLENGE", "What makes interviews difficult?"), ("ONE GOAL", "What evidence do you want to improve?"), ("PAIR UP", "You will rotate roles during activities.")])
    admin_cards("Ground Rules", [("RESPECT", "Protect role-play safety and confidentiality."), ("PARTICIPATE", "Practise, observe and give evidence-based feedback."), ("PRIVACY", "Use only de-identified training data in AI tools."), ("CHALLENGE", "Question AI output and document human decisions.")])
    admin_cards("Course Roadmap", [("TOPIC 1", TOPICS[0][1]), ("TOPIC 2", TOPICS[1][1]), ("TOPIC 3", TOPICS[2][1]), ("PRACTICE", "10 progressive activities and a capstone role play.")], accent=TEAL)
    admin_cards("Learning Outcomes", [("LO1", LEARNING_OUTCOMES[0]), ("LO2", LEARNING_OUTCOMES[1]), ("LO3", LEARNING_OUTCOMES[2]), ("EVIDENCE", "Guide, transcript, scoring and feedback artefacts.")], accent=BLUE)
    admin_cards("Skills Framework", [("TSC", f"{TSC_TITLE} | {TSC_CODE}"), ("KNOWLEDGE", "K1-K6: types, outcomes, constraints, objectives, questions, listening."), ("ABILITIES", "A1 manage | A2 deliver | A3 provide improvement input."), ("ALIGNMENT", "Every concept and activity carries K/A tags.")], accent=VIOLET)
    admin_cards("Briefing for Assessment", [("FORMAT", "Open book: course slides and Learner Guide."), ("INTEGRITY", "Individual responses; no discussion or recording."), ("SUBMISSION", "Complete the provided documents and upload to LMS."), ("SUPPORT", "Clarify administrative instructions before timing starts.")], accent=AMBER)
    admin_cards("Assessment and Funding", [("WA", "6 open-ended SAQs | K1-K6 | 30 minutes."), ("ROLE PLAY", "One scenario | A1-A3 | 30 minutes."), ("RESULT", "Competent / Not Yet Competent with appeal rights."), ("FUNDING", "At least 75 percent attendance and competent assessment.")], accent=VIOLET)
    flow_slide()
    lab_login_slide()
    copilot_env_slide()
    sharepoint_corpus_slide()
    prompt_pattern_slide()
    if (ASSETS / "copilot-chat-tgfep.png").exists():
        tool_screenshot_slide("Microsoft 365 Copilot Chat", ASSETS / "copilot-chat-tgfep.png", COPILOT_CHAT_URL,
                              ["Sign in with your training account.",
                               "Ask a job-related question in your own words.",
                               "Read the answer critically and verify legal claims at source.",
                               "Never paste identifiers or a real candidate CV."])
    if (ASSETS / "copilot-agents-list.png").exists():
        tool_screenshot_slide("Published Agents in Your Environment", ASSETS / "copilot-agents-list.png", COPILOT_AGENTS_URL,
                              ["Five reference agents, already published.",
                               "Open one and read its Instructions before using it.",
                               "Build your own in Activity 6 and publish it.",
                               "Workflows named DO NOT DELETE are shared class assets."])
    tool_screenshot_slide("HR Interview Question Generator", ASSETS / "hr-question-generator.png", HR_TOOL_URL,
                          ["Enter de-identified CV and JD evidence.", "Inspect matches, gaps and unknowns.", "Review 15-plus questions across eight categories.", "Approve a printable interview guide."])
    tool_screenshot_slide("AI Interview Practice Lab", ASSETS / "ai-interview-practice.png", AI_PRACTICE_URL,
                          ["Choose one of six role-specific rubrics.", "Practise 6-10 questions in Demo mode.", "Receive criterion-level feedback.", "Upgrade the weakest answer and retry."])
    admin_cards("Courseware and Activities Access", [("LMS", LMS_URL), ("ACTIVITIES", "Each activity has its own instruction and checklist PDF."), ("GITHUB", GITHUB_URL), ("DATA", "Use supplied de-identified files; never upload live candidate data.")], accent=TEAL)
    for topic_no, topic_title, topic_codes in TOPICS:
        slide_map["topics"][str(topic_no)] = {"title": topic_title, "start": len(prs.slides)+1}
        section_slide(topic_no, topic_title, topic_codes)
        for concept_index, item in enumerate([c for c in CONCEPTS if c["topic"] == topic_no]):
            variant = concept_index % 4
            process_slide(item, variant)
            if item["chart"]: chart_slide(item)
            else: control_slide(item, variant)
            case_slide(item)
        for activity in [a for a in ACTIVITIES if a["topic"] == topic_no]:
            activity_slides(activity)
        cards_slide(f"Topic {topic_no} evidence recap", [("MECHANISM", "Name the process and control points."), ("EVIDENCE", "Retain traceable artefacts."), ("FAILURE", "Recognise bias, drift and missing evidence."), ("PRACTICE", "Apply the activity checklist before moving on.")], f"TOPIC {topic_no} | RECAP", accent=PALETTE[(topic_no-1)%4], takeaway="Explain one human decision that changed or rejected an AI suggestion.")
        if topic_no < 3:
            cards_slide("Break and reset", [("SAVE", "Save only de-identified activity evidence."), ("RESET", "Clear pasted data and close AI sessions."), ("REFLECT", "Write one question or control to revisit."), ("RETURN", "Be ready to rotate role-play positions.")], "COURSE BREAK", accent=TEAL, takeaway="Breaks are not evidence-retention periods: clear live personal data from shared devices.")
        slide_map["topics"][str(topic_no)]["end"] = len(prs.slides)
    cards_slide("What You Can Now Do", [("PLAN", "Create a fair, job-related interview evidence contract."), ("GENERATE", "Use GenAI to draft questions under explicit constraints."), ("CONDUCT", "Interview and answer with structure, listening and neutral probes."), ("DECIDE", "Score, compare and provide feedback from evidence.")], "COURSE SUMMARY", accent=TEAL, takeaway="AI accelerates preparation and practice; people remain accountable for fair hiring decisions.")
    source_slides()
    lms_orientation_slide()
    practice_exam_slide()
    admin_cards("Assessment Reminder", [("WA", "30 minutes | six open-ended questions."), ("ROLE PLAY", "30 minutes | interviewer performance A1-A3."), ("OPEN BOOK", "Use approved course slides and Learner Guide."), ("SUBMIT", "Upload the candidate papers to the LMS.")], accent=VIOLET)
    flow_slide()
    attendance_slide(end=True)
    slide = new_slide(numbered=True)
    add_rect(slide, Inches(0.72), Inches(1.80), Inches(11.85), Inches(4.65), LIGHT)
    add_text(slide, Inches(1.20), Inches(2.35), Inches(10.85), Inches(0.40), "THANK YOU", 16, TEAL, True, PP_ALIGN.CENTER)
    add_text(slide, Inches(1.20), Inches(3.05), Inches(10.85), Inches(1.10), "Interview with evidence.\nUse AI with judgement.", 36, INK, True, PP_ALIGN.CENTER)
    add_link_text(slide, Inches(3.00), Inches(5.15), Inches(7.3), Inches(0.34), COURSE_URL, COURSE_URL, 12, BLUE)
    for idx, slide in enumerate(prs.slides):
        _transition(slide, "push" if idx in section_slide_ids else "fade", "med" if idx in section_slide_ids else "fast")
    out = CW / f"{SHORT_TITLE}-v{VERSION}.pptx"
    prs.save(out)
    (CW / "slide_map.json").write_text(json.dumps(slide_map, indent=2), encoding="utf-8")
    return out


def _set_cell_text(cell, text, bold=False, size=9.5, color="161B26"):
    cell.text = ""
    p = cell.paragraphs[0]; p.paragraph_format.space_after = DPt(0)
    r = p.add_run(str(text)); r.bold = bold; r.font.name = "Arial"; r.font.size = DPt(size); r.font.color.rgb = DRGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); tcPr.append(shd)


def _set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None: tcMar = OxmlElement("w:tcMar"); tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None: node = OxmlElement(f"w:{m}"); tcMar.append(node)
        node.set(qn("w:w"), str(v)); node.set(qn("w:type"), "dxa")


def _repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader"); repeat.set(qn("w:val"), "true"); tr_pr.append(repeat)


def _cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    flag = OxmlElement("w:cantSplit"); tr_pr.append(flag)


def _field(paragraph, instr):
    run = paragraph.add_run(); begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText"); text.set(qn("xml:space"), "preserve"); text.text = instr
    separate = OxmlElement("w:fldChar"); separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, text, separate, end])


def _hyperlink(paragraph, text, url):
    part = paragraph.part; rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink"); link.set(qn("r:id"), rid)
    run = OxmlElement("w:r"); rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color"); color.set(qn("w:val"), "1F6FEB"); rpr.append(color)
    underline = OxmlElement("w:u"); underline.set(qn("w:val"), "single"); rpr.append(underline)
    run.append(rpr); t = OxmlElement("w:t"); t.text = text; run.append(t); link.append(run); paragraph._p.append(link)


def _doc_styles(doc):
    sec = doc.sections[0]; sec.top_margin = DInches(0.75); sec.bottom_margin = DInches(0.72); sec.left_margin = DInches(0.82); sec.right_margin = DInches(0.82)
    normal = doc.styles["Normal"]; normal.font.name = "Arial"; normal.font.size = DPt(11)
    normal.paragraph_format.space_after = DPt(6); normal.paragraph_format.line_spacing = 1.15
    for style_name, size, color in (("Heading 1", 16, "1F6FEB"), ("Heading 2", 13, "108A73"), ("Heading 3", 11.5, "6D3FD2")):
        s = doc.styles[style_name]; s.font.name = "Arial"; s.font.size = DPt(size); s.font.bold = True; s.font.color.rgb = DRGBColor.from_string(color)
        s.paragraph_format.space_before = DPt(12); s.paragraph_format.space_after = DPt(6)
    footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run(f"{ORG} | {COURSE_CODE} | Page "); r.font.name = "Arial"; r.font.size = DPt(8)
    _field(footer, "PAGE"); footer.add_run(" of "); _field(footer, "NUMPAGES")
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields"); settings.append(update)
    update.set(qn("w:val"), "true")


def _cover(doc, instrument):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = DPt(20)
    run = p.add_run(ORG); run.bold=True; run.font.name="Arial"; run.font.size=DPt(13)
    p=doc.add_paragraph(f"UEN: {UEN}"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("")
    p=doc.add_paragraph(instrument.upper()); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs: r.bold=True; r.font.name="Arial"; r.font.size=DPt(25); r.font.color.rgb=DRGBColor.from_string("1F6FEB")
    p=doc.add_paragraph("For"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=doc.add_paragraph(TITLE); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs: r.bold=True; r.font.name="Arial"; r.font.size=DPt(20)
    p=doc.add_paragraph(f"TGS Ref No: {COURSE_CODE}"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("")
    p=doc.add_paragraph("Conducted by"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=doc.add_paragraph(ORG); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs: r.bold=True; r.font.name="Arial"; r.font.size=DPt(13)
    p=doc.add_paragraph(f"UEN: {UEN}"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=doc.add_paragraph(f"Version {VERSION}"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs: r.bold=True; r.font.name="Arial"; r.font.size=DPt(12)
    doc.add_page_break()


def _version_and_toc(doc, changes):
    doc.add_heading("DOCUMENT VERSION CONTROL RECORD", level=1)
    table = doc.add_table(rows=2, cols=4); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.autofit=False
    headers=["Version Number","Effective Date of Release","Summary of Included Changes","Author"]
    widths=[0.95,1.35,3.65,1.05]
    for j,h in enumerate(headers): _set_cell_text(table.cell(0,j),h,True,9,"FFFFFF"); _shade(table.cell(0,j),"1F6FEB"); table.cell(0,j).width=DInches(widths[j]); _set_cell_margins(table.cell(0,j))
    _repeat_header(table.rows[0]); _cant_split(table.rows[0])
    vals=[VERSION,VERSION_DATE,changes,"Dr Alfred Ang"]
    for j,v in enumerate(vals): _set_cell_text(table.cell(1,j),v,False,9); table.cell(1,j).width=DInches(widths[j]); _set_cell_margins(table.cell(1,j))
    doc.add_page_break(); doc.add_heading("TABLE OF CONTENTS", level=1)
    p=doc.add_paragraph(); _field(p, 'TOC \\o "1-3" \\h \\z \\u')
    doc.add_paragraph("Open in Microsoft Word and update the table if page numbers change.")
    doc.add_page_break()


def build_learner_guide_md():
    """Markdown mirror of the Learner Guide - same single source, so the two cannot diverge."""
    L = []
    L.append(f"# Learner Guide - {TITLE}")
    L.append("")
    L.append(f"**Course code:** {COURSE_CODE}  |  **Version:** {VERSION}  |  **Date:** {VERSION_DATE}")
    L.append(f"**Provider:** {ORG} (UEN {UEN})  |  **Trainer:** {TRAINER}")
    L.append("")
    L.append("## How to Use This Guide")
    L.append("Use the concept sections before each activity, then follow the detailed steps in the matching "
             "activity folder. Work only with de-identified training data. The slides explain mechanisms and "
             "decision rules; this guide contains the complete operational procedure.")
    L.append("")
    L.append(f"[Course LMS]({LMS_URL}) | [Microsoft 365 Copilot]({M365_PORTAL}) | "
             f"[Copilot Studio environment]({COPILOT_AGENTS_URL}) | [AI Interview Practice Lab]({AI_PRACTICE_URL})")
    L.append("")
    L.append("## Your Lab Sign-In")
    L.append(f"- Portal: {M365_PORTAL}")
    L.append(f"- Accounts: `{LAB_LOGIN_USER_1}` or `{LAB_LOGIN_USER_2}`")
    L.append(f"- **{LAB_LOGIN_NOTE}**")
    L.append("")
    L.append("### Published agents in the course environment")
    for an, ad, au in COPILOT_AGENTS:
        L.append(f"- **{an}** - {ad} ([open]({au}))")
    L.append("")
    L.append(f"### SharePoint practice corpus")
    L.append(SHAREPOINT_CORPUS)
    for sn, ss, su in SHAREPOINT_SITES:
        L.append(f"- [{sn}]({su}) - {ss}")
    L.append("")
    L.append("## Course Outcomes and Assessment")
    for lo in LEARNING_OUTCOMES: L.append(f"- {lo}")
    L.append("")
    L.append("Assessment: 30-minute Written Assessment (six open-ended SAQs covering K1-K6) followed by a "
             "30-minute Role Play observed against A1-A3.")
    L.append(f"Practice exam: {PRACTICE_EXAM_URL} - attempt it before the Written Assessment.")
    L.append("")
    for topic_no, topic_title, topic_codes in TOPICS:
        L.append(f"## Topic {topic_no}: {topic_title}")
        L.append(f"*Alignment: {topic_codes}*")
        L.append("")
        for item in [c for c in CONCEPTS if c["topic"] == topic_no]:
            L.append(f"### {item['title']}")
            L.append(item["rule"])
            L.append("")
            L.append("| Mechanism | Control / evidence |")
            L.append("|---|---|")
            for st, (head, body) in zip(item["mechanism"], item["controls"]):
                L.append(f"| {st} | **{head}:** {body} |")
            L.append("")
            L.append(f"**Worked evidence:** Input: {item['case'][0]}  AI risk: {item['case'][1]}  "
                     f"Human review: {item['case'][2]}  Evidence: {item['case'][3]}")
            L.append(f"**Source:** {item['source']}")
            L.append("")
        for a in [x for x in ACTIVITIES if x["topic"] == topic_no]:
            L.append(f"### Activity {a['num']}: {a['title']}")
            L.append(f"- **Goal:** {a['outcome']}")
            L.append(f"- **Scenario:** {a['scenario']}")
            L.append(f"- **Roles:** {a['roles']}")
            L.append(f"- **Tools:** {a['tools']}")
            L.append(f"- **Duration:** {a['duration']} minutes")
            L.append("")
            L.append("**Before you start.** Use only the supplied de-identified scenario and files. Keep the "
                     "instruction PDF and checklist PDF open from the activity folder. Do not enter live "
                     "candidate, employer-confidential or production API-key data.")
            L.append("")
            L.append("**Step-by-step**")
            for i, step in enumerate(a["steps"], 1): L.append(f"{i}. {step}")
            L.append("")
            L.append("**Evidence to save**")
            for name in a["files"]: L.append(f"- `{name}`")
            L.append("")
            L.append("**Acceptance checklist**")
            for check in a["checklist"]: L.append(f"- [ ] {check}")
            L.append("")
            L.append(f"Folder: `activities/activity-{a['num']:02d}-{slug(a['title'])}/`")
            L.append("")
    L.append("## Assessment Flow")
    for i, step in enumerate(["TRAQOM digital attendance", "Assessment digital attendance",
                              "Written Assessment then Role Play",
                              "Upload completed candidate papers to the LMS",
                              "Sign the Assessment Summary Record"], 1):
        L.append(f"{i}. {step}")
    L.append("")
    L.append("## Sources and Further Reading")
    for key, url in SOURCES.items():
        L.append(f"- **{key.replace('_',' ').title()}:** {url}")
    L.append("")
    L.append("---")
    L.append(f"(c) {ORG}. Synthetic training data only. AI output is a draft for human review.")
    out = CW / f"LG-{SHORT_TITLE}-v{VERSION}.md"
    out.write_text("\n".join(L), encoding="utf-8")
    return out


def _activity_image(activity):
    """Pick the most relevant course screenshot for an activity's Learner Guide section."""
    t = (activity["title"] + " " + activity["tools"]).lower()
    if "practice lab" in t or "web app" in t:
        return ASSETS / "ai-interview-practice.png"
    if "copilot studio" in t or "agent" in t:
        return ASSETS / "copilot-agents-list.png"
    if "question generator" in t:
        return ASSETS / "hr-question-generator.png"
    return ASSETS / "copilot-chat-tgfep.png"


def build_learner_guide():
    doc=Document(); _doc_styles(doc); _cover(doc,"Learner Guide")
    _version_and_toc(doc,"Rebuilt for both interviewer and interviewee perspectives; 10 AI-supported role-play activities; current fairness, privacy and human-oversight controls.")
    doc.add_heading("How to Use This Guide", level=1)
    doc.add_paragraph("Use the concept sections before each activity, then follow the detailed steps in the matching activity folder. Work only with de-identified training data. The slides explain mechanisms and decision rules; this guide contains the complete operational procedure.")
    p=doc.add_paragraph(); _hyperlink(p,"Course LMS",LMS_URL); p.add_run(" | "); _hyperlink(p,"Microsoft 365 Copilot",M365_PORTAL); p.add_run(" | "); _hyperlink(p,"Copilot Studio environment",COPILOT_AGENTS_URL); p.add_run(" | "); _hyperlink(p,"AI Interview Practice Lab",AI_PRACTICE_URL)
    doc.add_heading("Your Lab Sign-In", level=1)
    doc.add_paragraph(f"Portal: {M365_PORTAL}")
    doc.add_paragraph(f"Accounts: {LAB_LOGIN_USER_1} or {LAB_LOGIN_USER_2}")
    doc.add_paragraph(LAB_LOGIN_NOTE)
    doc.add_paragraph("Published agents in the course environment:")
    for _an,_ad,_au in COPILOT_AGENTS:
        _p=doc.add_paragraph(style="List Bullet"); _p.add_run(_an+" - ").bold=True; _p.add_run(_ad+"  "); _hyperlink(_p,"open",_au)
    doc.add_paragraph("SharePoint practice corpus: "+SHAREPOINT_CORPUS)
    for _sn,_ss,_su in SHAREPOINT_SITES:
        _p=doc.add_paragraph(style="List Bullet"); _p.add_run(f"{_sn} ({_ss}): "); _hyperlink(_p,_su,_su)
    doc.add_heading("Course Outcomes and Assessment", level=1)
    for lo in LEARNING_OUTCOMES: doc.add_paragraph(lo, style="List Bullet")
    doc.add_paragraph("Assessment: 30-minute Written Assessment (six open-ended SAQs covering K1-K6) followed by a 30-minute Role Play observed against A1-A3.")
    p=doc.add_paragraph(); p.add_run("Practice exam: ").bold=True; _hyperlink(p,PRACTICE_EXAM_URL,PRACTICE_EXAM_URL); p.add_run(" - attempt it before the Written Assessment.")
    for topic_no, topic_title, topic_codes in TOPICS:
        doc.add_heading(f"Topic {topic_no}: {topic_title}", level=1)
        doc.add_paragraph(f"Alignment: {topic_codes}")
        for item in [c for c in CONCEPTS if c["topic"]==topic_no]:
            doc.add_heading(item["title"], level=2)
            doc.add_paragraph(item["rule"])
            table=doc.add_table(rows=1, cols=2); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.autofit=False
            _set_cell_text(table.cell(0,0),"Mechanism",True,9,"FFFFFF"); _set_cell_text(table.cell(0,1),"Control / evidence",True,9,"FFFFFF"); _shade(table.cell(0,0),"1F6FEB"); _shade(table.cell(0,1),"1F6FEB")
            table.cell(0,0).width=DInches(2.05); table.cell(0,1).width=DInches(4.75)
            _repeat_header(table.rows[0]); _cant_split(table.rows[0])
            for st,(head,body) in zip(item["mechanism"],item["controls"]):
                row=table.add_row().cells; _set_cell_text(row[0],st,True,9); _set_cell_text(row[1],f"{head}: {body}",False,9)
                for c in row: _set_cell_margins(c)
                row[0].width=DInches(2.05); row[1].width=DInches(4.75); _cant_split(table.rows[-1])
            p=doc.add_paragraph(); p.add_run("Worked evidence: ").bold=True; p.add_run(f"Input: {item['case'][0]} AI risk: {item['case'][1]} Human review: {item['case'][2]} Evidence: {item['case'][3]}")
            p=doc.add_paragraph(); p.add_run("Source: ").bold=True; _hyperlink(p,item["source"],item["source"])
        for activity in [a for a in ACTIVITIES if a["topic"]==topic_no]:
            doc.add_page_break(); doc.add_heading(f"Activity {activity['num']}: {activity['title']}", level=2)
            doc.add_paragraph(f"Goal: {activity['outcome']}")
            doc.add_paragraph(f"Scenario: {activity['scenario']}")
            doc.add_paragraph(f"Roles: {activity['roles']}")
            doc.add_paragraph(f"Tools: {activity['tools']}")
            doc.add_heading("Before You Start", level=3)
            doc.add_paragraph("Use only the supplied de-identified scenario and files. Keep the instruction PDF and checklist PDF open from the activity folder. Do not enter live candidate, employer-confidential or production API-key data.")
            _img = _activity_image(activity)
            if _img and _img.exists():
                try:
                    doc.add_picture(str(_img), width=DInches(6.1))
                    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                    _cap = doc.add_paragraph(f"Figure {activity['num']}: reference screen for Activity {activity['num']}.")
                    _cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for _r in _cap.runs: _r.font.size = DPt(9); _r.italic = True
                except Exception:
                    pass
            doc.add_heading("Step-by-step", level=3)
            for i,step in enumerate(activity["steps"],1): doc.add_paragraph(f"{i}. {step}")
            doc.add_heading("Evidence to Save", level=3)
            for name in activity["files"]: doc.add_paragraph(name, style="List Bullet")
            doc.add_heading("Acceptance Checklist", level=3)
            for check in activity["checklist"]: doc.add_paragraph(f"[ ] {check}")
            doc.add_paragraph(f"Folder: activities/activity-{activity['num']:02d}-{slug(activity['title'])}/")
    doc.add_page_break(); doc.add_heading("Assessment Flow", level=1)
    for i,step in enumerate(["TRAQOM digital attendance","Assessment digital attendance","Written Assessment then Role Play","Upload completed candidate papers to the LMS","Sign the Assessment Summary Record"],1): doc.add_paragraph(f"{i}. {step}")
    doc.add_page_break(); doc.add_heading("Sources and Further Reading", level=1)
    for key,url in SOURCES.items():
        p=doc.add_paragraph(); p.add_run(key.replace("_"," ").title()+": ").bold=True; _hyperlink(p,url,url)
    out=CW/f"LG-{SHORT_TITLE}-v{VERSION}.docx"; doc.save(out); return out


def build_lesson_plan():
    doc=Document(); _doc_styles(doc); _cover(doc,"Lesson Plan")
    _version_and_toc(doc,"Rebuilt one-day plan for three topics, 10 aligned activities, WA plus Role Play, and exact slide references from slide_map.json.")
    doc.add_heading("Course Overview",level=1)
    doc.add_paragraph("A one-day, 8-hour WSQ programme for interviewers and interviewees to plan, conduct, evaluate and improve structured interviews with accountable generative-AI support.")
    doc.add_heading("Learning Outcomes",level=1)
    for lo in LEARNING_OUTCOMES: doc.add_paragraph(lo,style="List Bullet")
    doc.add_heading("Daily Schedule",level=1)
    rows=[
        ("9:30-9:45","15 min","Attendance, TRAQOM and course administration","Briefing",f"Slides 1-{slide_map['topics']['1']['start']-1}"),
        ("9:45-11:00","75 min","Topic 1 foundations and Activity 1","Concepts, sign-in, guided practice",f"Slides {slide_map['topics']['1']['start']}-{slide_map['activities']['1']['end']}"),
        ("11:00-11:15","15 min","Morning tea break (counted within instructional time)","Break","-"),
        ("11:15-12:30","75 min","Topic 1 controls and Activities 2-3","Case analysis, red-team activity",f"Slides {slide_map['activities']['2']['start']}-{slide_map['topics']['1']['end']}"),
        ("12:30-1:30","60 min","Lunch (not counted as instructional time)","Break","-"),
        ("1:30-2:50","80 min","Topic 2 prompt engineering and Activities 4-5","Demonstration, prompt practice",f"Slides {slide_map['topics']['2']['start']}-{slide_map['activities']['5']['end']}"),
        ("2:50-3:05","15 min","Afternoon tea break (counted within instructional time)","Break","-"),
        ("3:05-4:20","75 min","Topic 2 Activities 6-7","Agent build, candidate practice and role play",f"Slides {slide_map['activities']['6']['start']}-{slide_map['topics']['2']['end']}"),
        ("4:20-5:20","60 min","Topic 3 Activities 8-10","Scoring, calibration, feedback and capstone",f"Slides {slide_map['topics']['3']['start']}-{slide_map['topics']['3']['end']}"),
        ("5:20-5:30","10 min","Assessment briefing and digital attendance","Briefing","Closing admin slides"),
        ("5:30-6:00","30 min","Written Assessment (WA)","Open-book individual assessment","-"),
        ("6:00-6:30","30 min","Role Play (RP)","Observed performance assessment","-"),
    ]
    table=doc.add_table(rows=1,cols=5); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.autofit=False
    heads=["Time","Duration","Topic / Activity","Method","Slides"]
    widths=[1.0,0.75,2.55,1.55,1.25]
    for j,h in enumerate(heads): _set_cell_text(table.cell(0,j),h,True,9.5,"FFFFFF"); _shade(table.cell(0,j),"1F6FEB"); table.cell(0,j).width=DInches(widths[j]); _set_cell_margins(table.cell(0,j))
    _repeat_header(table.rows[0]); _cant_split(table.rows[0])
    for rowdata in rows:
        row=table.add_row().cells
        for j,v in enumerate(rowdata): _set_cell_text(row[j],v,False,9.5); row[j].width=DInches(widths[j]); _set_cell_margins(row[j])
        _cant_split(table.rows[-1])
    doc.add_page_break(); doc.add_heading("Topic-by-topic Breakdown",level=1)
    for topic_no,topic_title,codes in TOPICS:
        sm=slide_map["topics"][str(topic_no)]
        doc.add_heading(f"Topic {topic_no}: {topic_title} | Slides {sm['start']}-{sm['end']}",level=2)
        doc.add_paragraph(f"Alignment: {codes}. Trainer uses mechanism, control, case and evidence slides; detailed operational steps remain in the Learner Guide and activity PDFs.")
        for a in [x for x in ACTIVITIES if x["topic"]==topic_no]:
            am=slide_map["activities"][str(a["num"])]
            doc.add_heading(f"Activity {a['num']}: {a['title']} | Slides {am['start']}-{am['end']}",level=3)
            doc.add_paragraph(f"{a['duration']} minutes | {a['roles']} | Output: {a['outcome']}")
    doc.add_heading("Resources Required",level=1)
    for item in ["Laptop with modern browser","Course LMS access","HR Interview Question Generator (local mode)","AI Interview Practice Lab (Demo mode)","De-identified activity files","Instruction and checklist PDFs in each activity folder"]: doc.add_paragraph(item,style="List Bullet")
    doc.add_heading("Assessment",level=1)
    doc.add_paragraph("5:30-6:00 PM: Written Assessment, six open-ended SAQs covering K1-K6. 6:00-6:30 PM: Role Play observed against A1-A3. Both are open-book using approved course materials. Candidate papers are submitted through the LMS; answer keys and assessor checklist remain trainer-controlled.")
    p=doc.add_paragraph(); _hyperlink(p,LMS_URL,LMS_URL)
    out=CW/f"LP-{SHORT_TITLE}-v{VERSION}.docx"; doc.save(out); return out


def slug(text):
    return re.sub(r"[^a-z0-9]+","-",text.lower()).strip("-")


def pdf_header(canvas, doc, title):
    canvas.saveState(); canvas.setFillColor(colors.HexColor("#1F6FEB")); canvas.rect(0,A4[1]-14*mm,A4[0],14*mm,fill=1,stroke=0)
    canvas.setFillColor(colors.white); canvas.setFont("Helvetica-Bold",10); canvas.drawString(18*mm,A4[1]-9*mm,title)
    canvas.setFillColor(colors.HexColor("#5B6372")); canvas.setFont("Helvetica",8); canvas.drawString(18*mm,10*mm,f"{ORG} | {COURSE_CODE}")
    canvas.drawRightString(A4[0]-18*mm,10*mm,f"Page {doc.page}"); canvas.restoreState()


def normalize_pdf(path):
    """Rewrite page objects to avoid a Poppler clipping defect in some generated PDFs."""
    path = Path(path); tmp = path.with_suffix(".normalized.pdf")
    reader = PdfReader(str(path)); writer = PdfWriter()
    for page in reader.pages: writer.add_page(page)
    with tmp.open("wb") as handle: writer.write(handle)
    tmp.replace(path)


def build_activity_pdfs():
    styles=getSampleStyleSheet()
    styles.add(ParagraphStyle(name="ActTitle",parent=styles["Title"],fontName="Helvetica-Bold",fontSize=20,leading=24,textColor=colors.HexColor("#161B26"),spaceAfter=12))
    styles.add(ParagraphStyle(name="ActH",parent=styles["Heading2"],fontName="Helvetica-Bold",fontSize=12,leading=15,textColor=colors.HexColor("#1F6FEB"),spaceBefore=8,spaceAfter=5))
    styles.add(ParagraphStyle(name="ActBody",parent=styles["BodyText"],fontName="Helvetica",fontSize=9.5,leading=13,spaceAfter=5))
    styles.add(ParagraphStyle(name="ActSmall",parent=styles["BodyText"],fontName="Helvetica",fontSize=8,leading=10,textColor=colors.HexColor("#5B6372")))
    for a in ACTIVITIES:
        folder=ACT/f"activity-{a['num']:02d}-{slug(a['title'])}"; folder.mkdir(parents=True,exist_ok=True)
        for name,content in a["files"].items(): (folder/name).write_text(content,encoding="utf-8")
        (folder/"README.md").write_text(f"# Activity {a['num']}: {a['title']}\n\nUse `instruction.pdf` for the complete procedure and `checklist.pdf` for acceptance. Work only with the supplied de-identified files.\n",encoding="utf-8")
        inst=folder/"instruction.pdf"
        doc=BaseDocTemplate(str(inst),pagesize=A4,rightMargin=18*mm,leftMargin=18*mm,topMargin=22*mm,bottomMargin=18*mm)
        frame=Frame(doc.leftMargin,doc.bottomMargin,doc.width,doc.height,id="normal")
        doc.addPageTemplates([PageTemplate(id="activity",frames=[frame],onPage=lambda c,d,t=f"Activity {a['num']} Instruction":pdf_header(c,d,t))])
        story=[Spacer(1,8*mm),Paragraph(f"Activity {a['num']}: {a['title']}",styles["ActTitle"]),
               Paragraph(f"<b>Duration:</b> {a['duration']} minutes &nbsp;&nbsp; <b>Roles:</b> {a['roles']}",styles["ActBody"]),
               Paragraph("Scenario",styles["ActH"]),Paragraph(a["scenario"],styles["ActBody"]),
               Paragraph("Goal and output",styles["ActH"]),Paragraph(a["outcome"],styles["ActBody"]),
               Paragraph("Tools and data boundary",styles["ActH"]),Paragraph(a["tools"],styles["ActBody"]),
               Paragraph("Use only supplied de-identified training data. Do not enter live candidate data, employer-confidential information or production API keys.",styles["ActBody"]),
               Paragraph("Detailed step-by-step",styles["ActH"])]
        for i,step in enumerate(a["steps"],1): story.append(KeepTogether([Paragraph(f"<b>Step {i}</b>",styles["ActBody"]),Paragraph(step,styles["ActBody"])]))
        story.extend([Paragraph("Evidence files",styles["ActH"])])
        for name in a["files"]: story.append(Paragraph(f"- {name}",styles["ActBody"]));
        story.extend([Paragraph("Recovery and troubleshooting",styles["ActH"]),
                      Paragraph("If the tool output is missing, generic, biased or ungrounded, preserve the input and output, mark the item as rejected, revise one constraint at a time, and regenerate. If access fails, use the supplied templates and local/demo mode; do not substitute live candidate data.",styles["ActBody"]),
                      Paragraph("Completion criterion",styles["ActH"]),Paragraph(a["outcome"],styles["ActBody"])])
        doc.build(story)
        normalize_pdf(inst)
        chk=folder/"checklist.pdf"
        doc=BaseDocTemplate(str(chk),pagesize=A4,rightMargin=18*mm,leftMargin=18*mm,topMargin=22*mm,bottomMargin=18*mm)
        frame=Frame(doc.leftMargin,doc.bottomMargin,doc.width,doc.height,id="normal")
        doc.addPageTemplates([PageTemplate(id="check",frames=[frame],onPage=lambda c,d,t=f"Activity {a['num']} Checklist":pdf_header(c,d,t))])
        story=[Spacer(1,8*mm),Paragraph(f"Activity {a['num']} Acceptance Checklist",styles["ActTitle"]),
               Paragraph(a["title"],styles["ActH"]),Paragraph("Learner: ____________________  Date: __________  Observer: ____________________",styles["ActBody"])]
        data=[[Paragraph("Done",styles["ActSmall"]),Paragraph("Acceptance check",styles["ActSmall"]),Paragraph("Evidence / remarks",styles["ActSmall"])]]
        for check in a["checklist"]: data.append(["[  ]",Paragraph(check,styles["ActBody"]),""])
        table=Table(data,colWidths=[18*mm,92*mm,60*mm],rowHeights=[10*mm]+[18*mm]*len(a["checklist"]))
        table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#1F6FEB")),("TEXTCOLOR",(0,0),(-1,0),colors.white),("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),("VALIGN",(0,0),(-1,-1),"MIDDLE"),("ALIGN",(0,1),(0,-1),"CENTER"),("GRID",(0,0),(-1,-1),0.5,colors.HexColor("#D7E0EA")),("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#F5F8FC")]),("LEFTPADDING",(0,0),(-1,-1),6),("RIGHTPADDING",(0,0),(-1,-1),6),("TOPPADDING",(0,0),(-1,-1),5),("BOTTOMPADDING",(0,0),(-1,-1),5)]))
        story.extend([table,Spacer(1,8*mm),Paragraph("Overall: [  ] Ready  [  ] Revise and retry",styles["ActH"]),Paragraph("Observer signature: ______________________________",styles["ActBody"])])
        doc.build(story)
        normalize_pdf(chk)


def main():
    ppt=build_deck(); lg=build_learner_guide(); lgmd=build_learner_guide_md(); lp=build_lesson_plan(); build_activity_pdfs()
    print(ppt); print(lg); print(lgmd); print(lp); print(f"activities={len(ACTIVITIES)} pdfs={len(list(ACT.rglob('*.pdf')))} slides={len(prs.slides)}")


if __name__ == "__main__": main()
