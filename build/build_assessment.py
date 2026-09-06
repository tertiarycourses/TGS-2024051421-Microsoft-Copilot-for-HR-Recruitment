#!/usr/bin/env python3
"""Build the registered WA (SAQ) and Role Play assessment instruments."""

from pathlib import Path
import json

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from build_all import _doc_styles, _hyperlink, _set_cell_text, _shade, _set_cell_margins, _repeat_header, _cant_split
from course_content import *

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "assessment"
OUT.mkdir(exist_ok=True)
SLIDE_MAP = json.loads((ROOT / "courseware" / "slide_map.json").read_text(encoding="utf-8"))


def cover(doc, instrument, trainer=False):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(20)
    r = p.add_run(ORG); r.bold = True; r.font.name = "Arial"; r.font.size = Pt(13)
    p = doc.add_paragraph(f"UEN: {UEN}"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("")
    p = doc.add_paragraph(instrument.upper()); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs: r.bold = True; r.font.name = "Arial"; r.font.size = Pt(24); r.font.color.rgb = RGBColor(0x1F,0x6F,0xEB)
    p = doc.add_paragraph("For"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(TITLE); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs: r.bold = True; r.font.name = "Arial"; r.font.size = Pt(19)
    p = doc.add_paragraph(f"TGS Ref No: {COURSE_CODE}"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if trainer:
        p = doc.add_paragraph("TRAINER / ASSESSOR COPY - DO NOT DISTRIBUTE TO LEARNERS"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs: r.bold = True; r.font.color.rgb = RGBColor(0xC2,0x41,0x3A); r.font.size = Pt(11)
    doc.add_paragraph("")
    p = doc.add_paragraph(f"Conducted by\n{ORG}\nUEN: {UEN}\nVersion {VERSION}"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()


def candidate_information(doc, instrument, duration):
    doc.add_heading("Candidate Information and Instructions", level=1)
    table = doc.add_table(rows=4, cols=2); table.autofit = False
    fields = [("Candidate name", ""), ("NRIC / ID (last 4 characters)", ""), ("Assessment date", ""), ("Assessor", "")]
    for i,(label,value) in enumerate(fields):
        _set_cell_text(table.cell(i,0), label, True, 10); _set_cell_text(table.cell(i,1), value, False, 10)
        table.cell(i,0).width = Pt(185); table.cell(i,1).width = Pt(300)
        for cell in table.rows[i].cells: _set_cell_margins(cell, 110, 120, 110, 120)
        _cant_split(table.rows[i])
    doc.add_heading("Instructions", level=2)
    items = [
        f"Instrument: {instrument}. Duration: {duration} minutes.",
        "This is an open-book individual assessment using only the approved course slides and Learner Guide.",
        "Read every task before responding. Use job-related evidence and show the reasoning behind each decision.",
        "Do not discuss answers, record another participant, or paste assessment content into an external AI system.",
        "Submit the completed candidate paper through the course LMS when instructed.",
    ]
    for item in items: doc.add_paragraph(item, style="List Bullet")
    p = doc.add_paragraph(); p.add_run("Course LMS: ").bold = True; _hyperlink(p, LMS_URL, LMS_URL)
    doc.add_heading("Official Use Only", level=2)
    doc.add_paragraph("Result: [  ] Competent  [  ] Not Yet Competent     Appeal explained: [  ] Yes")
    doc.add_paragraph("Assessor comments: ______________________________________________________________________")
    doc.add_paragraph("Assessor signature: ______________________________  Date: __________________")
    doc.add_page_break()


def build_wa():
    doc = Document(); _doc_styles(doc); cover(doc, "Written Assessment (Short Answer Questions)")
    candidate_information(doc, "Written Assessment - Short Answer Questions (WA/SAQ)", 30)
    for idx, (code, question) in enumerate(WA_ITEMS, 1):
        if idx > 1 and idx % 2 == 1: doc.add_page_break()
        doc.add_heading(f"Question {idx} - {code}", level=2)
        doc.add_paragraph(question)
        doc.add_paragraph("Response:")
        for _ in range(5): doc.add_paragraph("________________________________________________________________________________")
    path = OUT / f"WA (SAQ) - {SHORT_TITLE} - v{VERSION}.docx"; doc.save(path); return path


def build_wa_key():
    doc = Document(); _doc_styles(doc); cover(doc, "Answer Key - Written Assessment (SAQ)", True)
    doc.add_heading("Assessor Guidance", level=1)
    doc.add_paragraph("Award competence only where the response demonstrates the required knowledge, applies it to the scenario and avoids unsupported conclusions. Equivalent technically sound wording is acceptable.")
    mapping = {"K3":"1", "K4":"1", "K1":"2", "K5":"2", "K2":"3", "K6":"3"}
    for idx, ((code, question), points) in enumerate(zip(WA_ITEMS, WA_ANSWERS), 1):
        doc.add_page_break()
        sm = SLIDE_MAP["topics"][mapping[code]]
        doc.add_heading(f"Question {idx} - {code} | Suggested Answer", level=1)
        doc.add_paragraph(question)
        doc.add_heading("Essential answer points", level=2)
        for point in points: doc.add_paragraph(point, style="List Bullet")
        doc.add_heading("Competence decision", level=2)
        doc.add_paragraph("Competent response: covers at least three relevant points, explains their application to the scenario, and contains no material legal, ethical or interviewing error.")
        doc.add_paragraph(f"Courseware reference: Topic {mapping[code]}, slides {sm['start']}-{sm['end']}; corresponding Learner Guide topic section.")
    path = OUT / f"Answer to WA (SAQ) - {SHORT_TITLE} - v{VERSION}.docx"; doc.save(path); return path


def build_role_play():
    doc = Document(); _doc_styles(doc); cover(doc, "Role Play Assessment")
    candidate_information(doc, "Role Play (RP)", 30)
    doc.add_heading("Role Play Scenario", level=1)
    doc.add_paragraph(RP_SCENARIO)
    doc.add_heading("Candidate Task", level=2)
    tasks = [
        "(A1) Review the de-identified Senior Data Analyst scenario and the human-approved GenAI-assisted interview guide.",
        "(A1) Brief Alex Lee on the purpose, structure, timing and evidence process; confirm access needs before questioning.",
        "(A2) Ask consistent job-related core questions, listen actively and use neutral probes to clarify personal action and results.",
        "(A3) Close professionally, record evidence separately from inference, score against the anchors and identify missing evidence.",
        "(A3) Provide concise, respectful and actionable feedback supported by the observed evidence.",
    ]
    for task in tasks: doc.add_paragraph(task, style="List Number")
    doc.add_paragraph("Assessed abilities: A1 manage a fair and inclusive interview | A2 deliver structured questions and neutral probes | A3 provide evidence-based feedback. The code beside each task is the ability it evidences.")
    doc.add_heading("Permitted Resources", level=2)
    doc.add_paragraph("Course slides, Learner Guide, supplied interview guide, evidence sheet and scoring rubric. No live candidate data and no external AI prompting during the timed role play.")
    doc.add_heading("Evidence to Submit", level=2)
    for item in ["Completed interview notes", "Completed anchored scoring sheet", "Candidate feedback record", "Candidate reflection on one improvement"]: doc.add_paragraph(item, style="List Bullet")
    path = OUT / f"ROLE PLAY (RP) - {SHORT_TITLE} - v{VERSION}.docx"; doc.save(path); return path


def build_role_play_checklist():
    doc = Document(); _doc_styles(doc); cover(doc, "Role Play - Assessor Observation Checklist", True)
    doc.styles["Normal"].font.size = Pt(9.5)
    doc.styles["Normal"].paragraph_format.space_after = Pt(3)
    doc.add_heading("Assessor Information", level=1)
    doc.add_paragraph("Candidate: __________________________  Assessor: __________________________  Date: ______________")
    doc.add_paragraph("Observe the complete role play. Mark each criterion Met only when directly supported by observable behaviour or retained evidence. Record the time or evidence reference.")
    table = doc.add_table(rows=1, cols=5); table.autofit = False
    headers = ["Code", "Observable criterion", "Met", "Not met", "Evidence / time"]
    widths = [0.55, 3.50, 0.60, 0.70, 1.55]
    for j,h in enumerate(headers):
        table.columns[j].width = Inches(widths[j])
        _set_cell_text(table.cell(0,j),h,True,8,"FFFFFF"); _shade(table.cell(0,j),"1F6FEB"); table.cell(0,j).width=Inches(widths[j]); _set_cell_margins(table.cell(0,j),60,80,60,80)
    _repeat_header(table.rows[0]); _cant_split(table.rows[0])
    for code, criterion in RP_CHECKS:
        row = table.add_row().cells
        values = [code, criterion, "[  ]", "[  ]", ""]
        for j,value in enumerate(values): _set_cell_text(row[j],value,j==0,8); row[j].width=Inches(widths[j]); _set_cell_margins(row[j],60,80,60,80)
        _cant_split(table.rows[-1])
    doc.add_heading("Overall Decision", level=1)
    doc.add_paragraph("[  ] Competent - all critical A1, A2 and A3 behaviours are demonstrated with sufficient evidence.")
    doc.add_paragraph("[  ] Not Yet Competent - one or more critical behaviours are missing or materially unsafe/unfair.")
    doc.add_paragraph("Feedback and required improvement: ______________________________________________________")
    doc.add_paragraph("Assessor signature: ______________________________  Date: __________________")
    doc.add_paragraph("Courseware references: Activities 5 and 10 for conduct; Activities 8 and 9 for scoring and feedback; Topic 3 slides 159-235.")
    path = OUT / f"ROLE PLAY (RP) - Assessor Observation Checklist - {SHORT_TITLE} - v{VERSION}.docx"; doc.save(path); return path


if __name__ == "__main__":
    for file in [build_wa(), build_wa_key(), build_role_play(), build_role_play_checklist()]: print(file)
