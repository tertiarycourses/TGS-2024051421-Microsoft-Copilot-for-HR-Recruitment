#!/usr/bin/env python3
"""Insert a visible static TOC cache while preserving the live Word TOC field."""
import re
import sys
from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph
from pypdf import PdfReader

docx_path, pdf_path = sys.argv[1], sys.argv[2]
maxlevel = int(sys.argv[3]) if len(sys.argv) > 3 else 2
doc = Document(docx_path)

def norm(value):
    return re.sub(r"[^a-z0-9]", "", value.lower())

heads = []
for paragraph in doc.paragraphs:
    if not paragraph.style.name.startswith("Heading"):
        continue
    try:
        level = int(paragraph.style.name.split()[-1])
    except ValueError:
        continue
    if level <= maxlevel and paragraph.text.strip() and paragraph.text.strip() != "TABLE OF CONTENTS":
        heads.append((level, paragraph.text.strip()))

pages = [norm(page.extract_text() or "") for page in PdfReader(pdf_path).pages]
entries = []
cursor = 0
for level, text in heads:
    key = norm(text)[:24]
    found = None
    for idx in range(cursor, len(pages)):
        if key and key in pages[idx]:
            found = idx + 1; cursor = idx; break
    if found is None:
        for idx, page in enumerate(pages):
            if key and key in page:
                found = idx + 1; break
    entries.append((level, text, found or 1))

toc_pages = max(1, (len(entries) + 42) // 43)
offset = toc_pages - 1
if offset:
    entries = [(level, text, page + offset) for level, text, page in entries]

placeholder = None
for paragraph in doc.paragraphs:
    xml = paragraph._p.xml
    if "TOC " in xml and "instrText" in xml:
        placeholder = paragraph; break
if placeholder is None:
    raise SystemExit("No live TOC field found")

anchor = placeholder._p
for level, text, page in entries:
    element = anchor.makeelement(qn("w:p"), {})
    anchor.addprevious(element)
    paragraph = Paragraph(element, placeholder._parent)
    paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(6.25), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    paragraph.paragraph_format.left_indent = Inches(0.28 if level >= 2 else 0)
    paragraph.paragraph_format.space_after = Pt(2.5)
    run = paragraph.add_run(f"{text}\t{page}")
    run.font.name = "Arial"; run.font.size = Pt(10.5 if level == 1 else 9.5)
    run.bold = level == 1; run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

doc.save(docx_path)
print(f"visible_toc entries={len(entries)} pages={toc_pages} file={docx_path}")
